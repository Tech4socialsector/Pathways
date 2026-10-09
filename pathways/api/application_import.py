# Copyright (c) 2026, NLSIU and contributors
# For license information, please see license.txt

"""Applications page > Import: add applications received outside the portal
(email, paper, an older system) from an Excel or CSV file. One row is one
application; the candidate is found by email or created. preview_import
checks every row without saving; run_import saves the rows that pass, each
on its own, so one bad row never blocks the rest. Candidates are emailed
only when asked (portal login, application received)."""

import base64
import csv
import io
import re
from datetime import date, datetime

import frappe
from frappe import _
from frappe.utils import cint, flt, getdate, today

from pathways.utils.application_form import mobile_error

MAX_ROWS = 500
MAX_BYTES = 5 * 1024 * 1024
# Imported applications start early in the process: later stages
# (interviews, offers) create their own records and must be done in the app.
IMPORT_STATUSES = ("Submitted", "Under Review", "Shortlisted")
ELIGIBILITY = ("Pending", "Eligible", "Not Eligible")
GENDERS = ("Female", "Male", "Prefer not to say", "Other")
CATEGORIES = ("General", "SC", "ST", "OBC", "PWD", "Other")
DEGREE_LEVELS = ("Undergraduate", "Postgraduate", "Doctoral")

# (header, key, required, help)
COLUMNS = (
	("Job Opening ID", "job_opening", True, "From the Job Openings sheet, e.g. PWY-JOB-2026-00002"),
	("Full Name", "full_name", True, ""),
	("Email", "email", True, "Finds the candidate if they already exist"),
	("Mobile Number", "mobile_number", True, "10 digits, or +country code"),
	("Date of Birth", "date_of_birth", False, "YYYY-MM-DD or DD-MM-YYYY"),
	("Gender", "gender", False, " / ".join(GENDERS)),
	("Category", "category", False, " / ".join(CATEGORIES)),
	("Address", "address", False, ""),
	("Application Date", "application_date", False, "Defaults to today"),
	("Status", "status", False, " / ".join(IMPORT_STATUSES) + " (default Submitted)"),
	("Eligibility", "eligibility_status", False, " / ".join(ELIGIBILITY) + " (default Pending)"),
	("Source", "source", False, "A Candidate Source from Master Setup"),
	("Current Salary", "current_salary", False, "Number"),
	("Expected Salary", "expected_salary", False, "Number"),
	("Overall Experience (Years)", "overall_experience_years", False, "Number"),
	("Relevant Experience (Years)", "relevant_experience_years", False, "Number"),
	("Notice Period", "notice_period", False, "e.g. 30 days"),
	("Highest Degree Level", "degree_level", False, " / ".join(DEGREE_LEVELS)),
	("Degree Name", "degree_name", False, "Needed with a degree level"),
	("Institution", "institution_name", False, ""),
	("Year of Graduation", "year_of_graduation", False, "e.g. 2018"),
	("Percentage / CGPA", "percentage_or_cgpa", False, ""),
	("Specialization", "specialization", False, ""),
	("Current Designation", "designation", False, ""),
	("Current Employer", "employer_name", False, "Needed with a designation"),
)
HEADER_KEYS = {re.sub(r"[^a-z0-9]", "", h.lower()): key for h, key, _req, _help in COLUMNS}


def _require_create():
	if frappe.db.get_value("User", frappe.session.user, "user_type") != "System User" or not frappe.has_permission("Application", "create"):
		frappe.throw(_("You are not allowed to import applications."), frappe.PermissionError)


# ------------------------------------------------------------ template


@frappe.whitelist()
def download_template():
	"""The Excel template: the columns with a help row, and the job
	openings to choose from."""
	_require_create()
	from openpyxl import Workbook
	from openpyxl.styles import Alignment, Font, PatternFill
	from openpyxl.utils import get_column_letter
	from openpyxl.worksheet.datavalidation import DataValidation

	wb = Workbook()
	ws = wb.active
	ws.title = "Applications"
	head_fill = PatternFill("solid", fgColor="F2E6E8")
	for c, (header, _key, required, help_text) in enumerate(COLUMNS, start=1):
		cell = ws.cell(1, c, header + (" *" if required else ""))
		cell.font, cell.fill = Font(bold=True), head_fill
		note = ws.cell(2, c, help_text)
		note.font = Font(italic=True, color="777777", size=9)
		note.alignment = Alignment(wrap_text=True, vertical="top")
		ws.column_dimensions[get_column_letter(c)].width = max(16, min(34, len(header) + 6))
	ws.row_dimensions[2].height = 42
	ws.freeze_panes = "A3"

	jobs = frappe.get_all("Job Opening", filters={"status": ["not in", ["Draft", "Cancelled"]]}, fields=["name", "job_title", "position", "status"], order_by="creation desc")
	js = wb.create_sheet("Job Openings")
	js.append(["Job Opening ID", "Job Title", "Job Code", "Status"])
	for cell in js[1]:
		cell.font, cell.fill = Font(bold=True), head_fill
	for j in jobs:
		js.append([j.name, j.job_title, j.position or "", j.status])
	for col, width in zip("ABCD", (24, 50, 18, 16)):
		js.column_dimensions[col].width = width

	def choose(values, column, sheet_range=None):
		dv = DataValidation(type="list", formula1=sheet_range or '"' + ",".join(values) + '"', allow_blank=True)
		dv.add(f"{column}3:{column}{MAX_ROWS + 2}")
		ws.add_data_validation(dv)

	letter = {key: get_column_letter(i) for i, (_h, key, _r, _t) in enumerate(COLUMNS, start=1)}
	if jobs:
		choose(None, letter["job_opening"], f"'Job Openings'!$A$2:$A${len(jobs) + 1}")
	choose(GENDERS, letter["gender"])
	choose(CATEGORIES, letter["category"])
	choose(IMPORT_STATUSES, letter["status"])
	choose(ELIGIBILITY, letter["eligibility_status"])
	choose(DEGREE_LEVELS, letter["degree_level"])

	buffer = io.BytesIO()
	wb.save(buffer)
	frappe.local.response.filename = "Applications Import Template.xlsx"
	frappe.local.response.filecontent = buffer.getvalue()
	frappe.local.response.type = "download"


# ------------------------------------------------------------ reading


def _read_rows(filename, content):
	"""[(sheet row number, {key: value})] from an .xlsx or .csv file."""
	try:
		raw = base64.b64decode((content or "").split(",", 1)[-1])
	except Exception:
		frappe.throw(_("Could not read the file. Upload the Excel template (.xlsx) or a .csv file."))
	if not raw:
		frappe.throw(_("The file is empty."))
	if len(raw) > MAX_BYTES:
		frappe.throw(_("The file is larger than 5 MB."))
	name = (filename or "").lower()
	if name.endswith(".xlsx"):
		from openpyxl import load_workbook

		try:
			wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
		except Exception:
			frappe.throw(_("Could not open the Excel file. Save it as .xlsx and try again."))
		ws = wb["Applications"] if "Applications" in wb.sheetnames else wb.worksheets[0]
		table = [list(r) for r in ws.iter_rows(values_only=True)]
	elif name.endswith(".csv"):
		text = raw.decode("utf-8-sig", errors="replace")
		table = list(csv.reader(io.StringIO(text)))
	else:
		frappe.throw(_("Upload the Excel template (.xlsx) or a .csv file."))

	if not table:
		frappe.throw(_("The file is empty."))
	headers = [HEADER_KEYS.get(re.sub(r"[^a-z0-9]", "", str(h or "").lower())) for h in table[0]]
	missing = [h for h, key, required, _t in COLUMNS if required and key not in headers]
	if missing:
		frappe.throw(_("These columns are missing: {0}. Download the template and keep its header row.").format(", ".join(missing)))

	help_row = {re.sub(r"\s+", " ", t).strip() for _h, _k, _r, t in COLUMNS if t}
	rows = []
	for number, values in enumerate(table[1:], start=2):
		record = {key: values[i] for i, key in enumerate(headers) if key and i < len(values)}
		cells = [v for v in record.values() if v not in (None, "") and str(v).strip()]
		if not cells:
			continue
		# The template's grey help row.
		if all(re.sub(r"\s+", " ", str(v)).strip() in help_row for v in cells):
			continue
		rows.append((number, record))
	if len(rows) > MAX_ROWS:
		frappe.throw(_("The file has {0} rows; import at most {1} at a time.").format(len(rows), MAX_ROWS))
	return rows


def _text(value):
	if value is None:
		return ""
	if isinstance(value, float) and value.is_integer():
		value = int(value)
	if isinstance(value, (datetime, date)):
		return value.strftime("%Y-%m-%d")
	return re.sub(r"\s+", " ", str(value)).strip()


def _date(value, label, errors):
	if value in (None, ""):
		return None
	if isinstance(value, datetime):
		return value.date()
	if isinstance(value, date):
		return value
	text = _text(value)
	for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%d %b %Y", "%d %B %Y"):
		try:
			return datetime.strptime(text, fmt).date()
		except ValueError:
			pass
	errors.append(_("{0}: use YYYY-MM-DD or DD-MM-YYYY").format(label))
	return None


def _number(value, label, errors):
	if value in (None, ""):
		return None
	try:
		number = float(str(value).replace(",", "").strip())
	except ValueError:
		errors.append(_("{0}: enter a number").format(label))
		return None
	if number < 0:
		errors.append(_("{0}: cannot be negative").format(label))
		return None
	return number


def _choice(value, options, label, errors, default=None):
	text = _text(value)
	if not text:
		return default
	match = next((o for o in options if o.lower() == text.lower()), None)
	if not match:
		errors.append(_("{0}: must be one of {1}").format(label, ", ".join(options)))
	return match or default


def _check_rows(rows):
	"""Validate every row; [{row, values, errors, candidate, job_title}]."""
	jobs = {j.name: j for j in frappe.get_all("Job Opening", fields=["name", "job_title", "status", "track"])}
	sources = {s.lower(): s for s in frappe.get_all("Candidate Source", pluck="name")}
	emails = [_text(r.get("email")).lower() for _n, r in rows]
	candidates = {
		c.email.lower(): c for c in frappe.get_all("Candidate", filters={"email": ["in", [e for e in emails if e] or [""]]}, fields=["name", "email", "full_name"])
	}
	seen = {}
	out = []
	for number, r in rows:
		errors = []
		v = {}
		job = _text(r.get("job_opening"))
		if not job:
			errors.append(_("Job Opening ID is required"))
		elif job not in jobs:
			errors.append(_("Job Opening {0} does not exist").format(job))
		elif jobs[job].status in ("Draft", "Cancelled"):
			errors.append(_("Job Opening {0} is {1}").format(job, jobs[job].status))
		v["job_opening"] = job

		v["full_name"] = _text(r.get("full_name"))
		if not v["full_name"]:
			errors.append(_("Full Name is required"))
		email = _text(r.get("email")).lower()
		if not email:
			errors.append(_("Email is required"))
		elif not frappe.utils.validate_email_address(email):
			errors.append(_("{0} is not a valid email").format(email))
		elif frappe.db.get_value("User", email, "user_type") == "System User":
			errors.append(_("{0} is a staff login, not a candidate").format(email))
		v["email"] = email
		mobile = _text(r.get("mobile_number"))
		if not mobile:
			errors.append(_("Mobile Number is required"))
		elif mobile_error(mobile):
			errors.append(_("Mobile Number: {0}").format(mobile_error(mobile)))
		v["mobile_number"] = mobile

		v["date_of_birth"] = _date(r.get("date_of_birth"), _("Date of Birth"), errors)
		v["gender"] = _choice(r.get("gender"), GENDERS, _("Gender"), errors)
		v["category"] = _choice(r.get("category"), CATEGORIES, _("Category"), errors)
		v["address"] = _text(r.get("address"))
		applied = _date(r.get("application_date"), _("Application Date"), errors)
		if applied and applied > getdate(today()):
			errors.append(_("Application Date cannot be in the future"))
		v["application_date"] = applied or getdate(today())
		v["status"] = _choice(r.get("status"), IMPORT_STATUSES, _("Status"), errors, "Submitted")
		v["eligibility_status"] = _choice(r.get("eligibility_status"), ELIGIBILITY, _("Eligibility"), errors, "Pending")
		source = _text(r.get("source"))
		v["source"] = sources.get(source.lower()) if source else None
		if source and not v["source"]:
			errors.append(_("Source {0} is not in Master Setup > Candidate Sources").format(source))
		for key, label in (
			("current_salary", _("Current Salary")),
			("expected_salary", _("Expected Salary")),
			("overall_experience_years", _("Overall Experience")),
			("relevant_experience_years", _("Relevant Experience")),
		):
			v[key] = _number(r.get(key), label, errors)
		v["notice_period"] = _text(r.get("notice_period"))

		v["degree_level"] = _choice(r.get("degree_level"), DEGREE_LEVELS, _("Highest Degree Level"), errors)
		v["degree_name"] = _text(r.get("degree_name"))
		v["institution_name"] = _text(r.get("institution_name"))
		v["percentage_or_cgpa"] = _text(r.get("percentage_or_cgpa"))
		v["specialization"] = _text(r.get("specialization"))
		year = _text(r.get("year_of_graduation"))
		v["year_of_graduation"] = cint(year) if year else None
		if year and not (year.isdigit() and 1950 <= cint(year) <= getdate(today()).year + 1):
			errors.append(_("Year of Graduation: enter a year like 2018"))
		has_degree = any(v[k] for k in ("degree_name", "institution_name", "percentage_or_cgpa", "specialization", "year_of_graduation"))
		if (v["degree_level"] or has_degree) and not (v["degree_level"] and v["degree_name"]):
			errors.append(_("Give both Highest Degree Level and Degree Name, or leave the qualification columns empty"))
		v["designation"] = _text(r.get("designation"))
		v["employer_name"] = _text(r.get("employer_name"))
		if bool(v["designation"]) != bool(v["employer_name"]):
			errors.append(_("Give both Current Designation and Current Employer, or neither"))

		existing = candidates.get(email)
		if email and job:
			key = (email, job)
			if key in seen:
				errors.append(_("Same candidate and job as row {0}").format(seen[key]))
			else:
				seen[key] = number
			if existing and frappe.db.exists("Application", {"candidate": existing.name, "job_opening": job, "status": ["!=", "Withdrawn"]}):
				errors.append(_("{0} has already applied for this job").format(email))
		out.append(
			{
				"row": number,
				"values": v,
				"errors": errors,
				"candidate": existing.name if existing else "",
				"candidate_status": "Existing candidate" if existing else "New candidate",
				"job_title": jobs[job].job_title if job in jobs else "",
				"track": jobs[job].track if job in jobs else None,
			}
		)
	return out


def _summary(checked):
	return [
		{
			"row": c["row"],
			"full_name": c["values"]["full_name"],
			"email": c["values"]["email"],
			"job_opening": c["values"]["job_opening"],
			"job_title": c["job_title"],
			"status": c["values"]["status"],
			"candidate_status": c["candidate_status"],
			"errors": c["errors"],
		}
		for c in checked
	]


@frappe.whitelist(methods=["POST"])
def preview_import(filename, content):
	"""Check the file without saving anything."""
	_require_create()
	checked = _check_rows(_read_rows(filename, content))
	return {
		"rows": _summary(checked),
		"total": len(checked),
		"valid": sum(1 for c in checked if not c["errors"]),
		"new_candidates": len({c["values"]["email"] for c in checked if not c["errors"] and not c["candidate"]}),
	}


# ------------------------------------------------------------ importing


def _create(c, created_candidates):
	v = c["values"]
	candidate = c["candidate"] or created_candidates.get(v["email"])
	if not candidate:
		# An existing candidate's profile is never overwritten by an import.
		doc = frappe.new_doc("Candidate")
		doc.update(
			{
				"full_name": v["full_name"],
				"email": v["email"],
				"mobile_number": v["mobile_number"],
				"date_of_birth": v["date_of_birth"],
				"gender": v["gender"],
				"category": v["category"],
				"address": v["address"],
			}
		)
		doc.insert(ignore_permissions=True)
		candidate = doc.name
		created_candidates[v["email"]] = candidate
	app = frappe.new_doc("Application")
	app.update(
		{
			"candidate": candidate,
			"job_opening": v["job_opening"],
			"track": c["track"],
			"status": v["status"],
			"eligibility_status": v["eligibility_status"],
			"application_date": v["application_date"],
			"source": v["source"],
			"current_salary": flt(v["current_salary"]) if v["current_salary"] is not None else None,
			"expected_salary": flt(v["expected_salary"]) if v["expected_salary"] is not None else None,
			"overall_experience_years": v["overall_experience_years"],
			"relevant_experience_years": v["relevant_experience_years"],
			"notice_period": v["notice_period"],
		}
	)
	if v["eligibility_status"] != "Pending":
		app.eligibility_checked_by = frappe.session.user
		app.eligibility_checked_on = frappe.utils.now_datetime()
	if v["degree_level"]:
		app.append(
			"qualifications",
			{
				"degree_level": v["degree_level"],
				"degree_name": v["degree_name"],
				"other_institution": v["institution_name"],
				"year_of_graduation": v["year_of_graduation"],
				"percentage_or_cgpa": v["percentage_or_cgpa"],
				"specialization": v["specialization"],
			},
		)
	if v["designation"]:
		app.append("employment_history", {"designation": v["designation"], "employer_name": v["employer_name"], "is_current": 1})
	app.insert(ignore_permissions=True)
	app.add_comment("Info", _("Imported from a file by {0}.").format(frappe.utils.get_fullname(frappe.session.user)))
	return candidate, app


@frappe.whitelist(methods=["POST"])
def run_import(filename, content, send_login=0, send_acknowledgement=0):
	"""Save every row that passes the checks; each row on its own."""
	_require_create()
	from pathways.utils.candidate_account import claim_records, queue_portal_account

	checked = _check_rows(_read_rows(filename, content))
	created_candidates = {}
	done, failed = [], []
	for c in checked:
		if c["errors"]:
			failed.append({"row": c["row"], "full_name": c["values"]["full_name"], "error": "; ".join(c["errors"])})
			continue
		frappe.db.savepoint("import_row")
		try:
			new_candidate = not c["candidate"] and c["values"]["email"] not in created_candidates
			candidate, app = _create(c, created_candidates)
		except Exception as e:
			frappe.db.rollback(save_point="import_row")
			frappe.clear_last_message()
			message = str(e) if isinstance(e, frappe.ValidationError) else _("Could not save this row.")
			if not isinstance(e, frappe.ValidationError):
				frappe.log_error(title="Pathways application import failed", message=frappe.get_traceback())
			failed.append({"row": c["row"], "full_name": c["values"]["full_name"], "error": re.sub(r"<[^>]+>", "", message)})
			continue
		email = c["values"]["email"]
		if frappe.db.exists("User", email):
			claim_records(candidate, email)
		elif new_candidate and cint(send_login):
			queue_portal_account(candidate)
		if cint(send_acknowledgement):
			frappe.enqueue("pathways.api.application.send_acknowledgement", application_name=app.name, enqueue_after_commit=True)
		done.append({"row": c["row"], "full_name": c["values"]["full_name"], "application_id": app.application_id, "name": app.name})
	return {"done": done, "failed": failed}
