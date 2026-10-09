# Copyright (c) 2026, NLSIU and contributors
# For license information, please see license.txt

"""Interview panel (workflow folder "6. Final Interview"): each member of
the job's Selection Committee scores every candidate on the Interview
Assessment Form; the scores are averaged in the consolidated score sheet."""

import frappe
from frappe import _
from frappe.utils import cint, flt

from pathways.api.interviews import _can_manage
from pathways.pathways.doctype.interview_assessment.interview_assessment import interview_rubric

VERDICTS = ("Recommended", "Waitlist", "Not Recommended")
# Decided applications: their scores are part of the record and stay as they are.
DECIDED = ("Selected", "Not Selected", "Offer Extended", "Offer Accepted", "Offer Declined", "Documents Pending", "Documents Verified", "Joined", "Withdrawn")
# The Admin and Faculty forms ask for the candidate's area of specialization.
SPECIALIZATION_TRACKS = ("Admin", "Faculty")


def _my_committees(user=None):
	return frappe.get_all(
		"Selection Committee", filters=[["Committee Member Row", "member", "=", user or frappe.session.user]], pluck="name"
	)


def _panel_members(committee):
	if not committee:
		return []
	return frappe.get_all(
		"Committee Member Row", filters={"parent": committee, "parenttype": "Selection Committee"}, fields=["member"], order_by="idx asc", pluck="member"
	)


def score_summary(interviews):
	"""{interview: {scored, panel_size, average, max_score, percent}} for the
	recruitment team's lists."""
	if not interviews:
		return {}
	rows = frappe.get_all(
		"Interview Assessment", filters={"interview": ["in", list(interviews)]}, fields=["interview", "total_score", "max_score"]
	)
	committees = dict(frappe.get_all("Interview", filters={"name": ["in", list(interviews)]}, fields=["name", "selection_committee"], as_list=True))
	out = {}
	for name in interviews:
		mine = [r for r in rows if r.interview == name]
		avg = sum(flt(r.total_score) for r in mine) / len(mine) if mine else 0
		top = flt(mine[0].max_score) if mine else 0
		out[name] = {
			"scored": len(mine),
			"panel_size": len(_panel_members(committees.get(name))),
			"average": round(avg, 2),
			"max_score": top,
			"percent": round(avg * 100 / top, 1) if top else 0,
		}
	return out


@frappe.whitelist()
def get_my_panel():
	"""Interview Panel page: final interviews of the committees the user
	sits on (the recruitment team sees all of them), with the user's own
	score and how many panellists have scored."""
	manage = _can_manage()
	committees = _my_committees()
	if not committees and not manage:
		return {"interviews": [], "can_manage": 0, "is_panellist": 0}
	filters = {"round_type": "Final", "status": ["!=", "Cancelled"]}
	if not manage:
		filters["selection_committee"] = ["in", committees]
	rows = frappe.get_all(
		"Interview",
		filters=filters,
		fields=[
			"name", "application", "candidate_name", "job_title", "position", "status", "scheduled_datetime", "mode",
			"meeting_platform", "meeting_link", "location", "selection_committee",
			"application.application_id as application_id", "application.job_opening as job_opening", "application.status as application_status",
		],
		order_by="scheduled_datetime asc",
	)
	mine = {
		r.interview: r
		for r in frappe.get_all(
			"Interview Assessment",
			filters={"panelist": frappe.session.user, "interview": ["in", [r.name for r in rows] or [""]]},
			fields=["interview", "total_score", "max_score", "score_percent", "verdict", "modified"],
		)
	}
	summary = score_summary([r.name for r in rows])
	for r in rows:
		r.is_panellist = 1 if r.selection_committee in committees else 0
		r.my_score = mine.get(r.name)
		r.scores = summary.get(r.name)
	return {"interviews": rows, "can_manage": 1 if manage else 0, "is_panellist": 1 if committees else 0}


@frappe.whitelist()
def get_assessment_form(interview, panelist=None):
	"""The Interview Assessment Form for one candidate: the rubric with its
	questions and the panellist's scores so far. A panellist fills their own
	form; the recruitment team can type in any panellist's (signed paper)
	form, chosen with `panelist`."""
	iv = frappe.get_doc("Interview", interview)
	iv.check_permission("read")
	user = frappe.session.user
	members = _panel_members(iv.selection_committee)
	manage = _can_manage()
	if panelist and panelist != user and not (manage and panelist in members):
		frappe.throw(_("You can only enter your own scores."), frappe.PermissionError)
	if not panelist:
		if user in members or not manage:
			panelist = user
		else:
			# The first panellist still to score, for the recruitment team.
			done = set(frappe.get_all("Interview Assessment", filters={"interview": iv.name}, pluck="panelist"))
			panelist = next((m for m in members if m not in done), members[0] if members else user)
	if iv.round_type != "Final":
		frappe.throw(_("Only final interviews are scored by the panel."))
	app = frappe.db.get_value("Application", iv.application, ["name", "application_id", "job_opening", "status", "selection_outcome"], as_dict=True)
	track = frappe.db.get_value("Job Opening", app.job_opening, "track")
	rubric = interview_rubric(iv.name)
	if not rubric:
		frappe.throw(_("No active Final Interview rubric for the {0} track. Add one in Master Setup > Scoring Rubrics.").format(track or "?"))

	existing = frappe.db.get_value("Interview Assessment", {"interview": iv.name, "panelist": panelist}, "name")
	mine = frappe.get_doc("Interview Assessment", existing) if existing else None
	given = {r.criterion_label: r.score_given for r in (mine.criteria if mine else [])}

	is_panellist = user in members
	on_behalf = panelist != user
	locked = ""
	if iv.status == "Cancelled":
		locked = _("This interview was cancelled.")
	elif app.status in DECIDED or app.selection_outcome:
		locked = _("A decision has been recorded for this candidate ({0}), so the scores can no longer be changed.").format(app.selection_outcome or app.status)
	elif not members:
		locked = _("This job has no Selection Committee yet.")
	elif panelist not in members:
		locked = _("Only members of this job's Selection Committee can score.")

	return {
		"interview": {
			"name": iv.name, "candidate_name": iv.candidate_name, "candidate": iv.candidate, "job_title": iv.job_title, "position": iv.position,
			"scheduled_datetime": iv.scheduled_datetime, "status": iv.status, "mode": iv.mode, "meeting_platform": iv.get("meeting_platform"),
			"meeting_link": iv.meeting_link, "location": iv.get("location"),
		},
		"application": app,
		"brief": candidate_brief(iv.application),
		"track": track,
		"rubric": {
			"name": rubric.name,
			"pass_percent": flt(rubric.pass_threshold_percent) or 50,
			"max_score": sum(flt(c.max_score) for c in rubric.criteria),
			"criteria": [{"label": c.criterion_label, "max_score": flt(c.max_score), "guidance": c.guidance_text or "", "score": given.get(c.criterion_label)} for c in rubric.criteria],
		},
		"ask_specialization": 1 if track in SPECIALIZATION_TRACKS else 0,
		"mine": {
			"name": mine.name, "total_score": mine.total_score, "verdict": mine.verdict, "additional_comments": mine.recommendation,
			"area_of_specialization": mine.area_of_specialization, "modified": mine.modified,
		} if mine else None,
		"can_score": 0 if locked else 1,
		"locked_reason": locked,
		"panelist": panelist,
		"panellist_name": frappe.utils.get_fullname(panelist),
		"on_behalf": 1 if on_behalf else 0,
		"entered_by": frappe.utils.get_fullname(mine.entered_by) if mine and mine.get("entered_by") else "",
		# The recruitment team picks whose form to type in.
		"panel": _panel_status(iv.name, members) if manage else [],
		"is_panellist": 1 if is_panellist else 0,
	}


def _panel_status(interview, members):
	scored = {
		r.panelist: r
		for r in frappe.get_all("Interview Assessment", filters={"interview": interview}, fields=["panelist", "total_score", "max_score"])
	}
	return [
		{
			"user": m,
			"full_name": frappe.utils.get_fullname(m),
			"total_score": scored[m].total_score if m in scored else None,
			"max_score": scored[m].max_score if m in scored else None,
		}
		for m in members
	]


@frappe.whitelist(methods=["POST"])
def submit_assessment(interview, scores, verdict=None, additional_comments=None, area_of_specialization=None, panelist=None):
	"""Save an Interview Assessment Form for one candidate: the signed-in
	panellist's own, or (recruitment team) a panellist's paper form typed
	in, recorded with who entered it. Can be changed until a decision is
	recorded for the candidate."""
	form = get_assessment_form(interview, panelist)
	if not form["can_score"]:
		frappe.throw(form["locked_reason"], frappe.PermissionError)
	scores = frappe.parse_json(scores) if isinstance(scores, str) else (scores or {})
	if verdict and verdict not in VERDICTS:
		frappe.throw(_("Choose a recommendation."))

	rows = []
	for c in form["rubric"]["criteria"]:
		value = scores.get(c["label"])
		if value in (None, ""):
			frappe.throw(_("Enter a score for {0}.").format(c["label"]))
		value = flt(value)
		if value < 0 or value > c["max_score"]:
			frappe.throw(_("{0}: the score must be between 0 and {1}.").format(c["label"], f"{c['max_score']:g}"))
		rows.append({"criterion_label": c["label"], "max_score": c["max_score"], "score_given": value})

	name = (form["mine"] or {}).get("name")
	doc = frappe.get_doc("Interview Assessment", name) if name else frappe.new_doc("Interview Assessment")
	doc.update(
		{
			"interview": interview,
			"panelist": form["panelist"],
			"entered_by": frappe.session.user if form["on_behalf"] else None,
			"verdict": verdict or None,
			"recommendation": (additional_comments or "").strip() or None,
			"area_of_specialization": (area_of_specialization or "").strip() or None,
			"panelist_signature_confirmed": 1,
		}
	)
	doc.set("criteria", rows)
	# Access was checked above: a member of this interview's committee.
	doc.flags.ignore_permissions = True
	doc.save()
	return {"name": doc.name, "total_score": doc.total_score, "max_score": doc.max_score, "score_percent": doc.score_percent}


# ------------------------------------------------- consolidated score sheet

OUTCOMES = ("Selected", "Waitlisted", "Not Selected")
# Statuses from which a selection outcome can be set or changed: before any offer.
OUTCOME_FROM = ("Interview Scheduled", "Interview Completed", "Selected", "Not Selected")


def _can_view_sheet(job_opening):
	if _can_manage():
		return True
	return bool(frappe.db.exists("Selection Committee", [["Committee Member Row", "member", "=", frappe.session.user], ["job_opening", "=", job_opening]]))


@frappe.whitelist()
def get_consolidated(job_opening, interview_date=None):
	"""Workflow folder "6. Final Interview / Consolidated score sheet": per
	candidate, every panellist's scores, the total of all and the average,
	against the 50% bar; with each panellist's recommendation."""
	if not _can_view_sheet(job_opening):
		frappe.throw(_("Not permitted."), frappe.PermissionError)
	job = frappe.db.get_value("Job Opening", job_opening, ["name", "job_title", "position", "track"], as_dict=True)
	committee = frappe.db.get_value("Selection Committee", {"job_opening": job_opening}, ["name", "scheduled_date"], as_dict=True)
	members = _panel_members(committee.name) if committee else []
	interviews = frappe.get_all(
		"Interview",
		filters={"round_type": "Final", "status": ["!=", "Cancelled"], "application": ["in", frappe.get_all("Application", filters={"job_opening": job_opening}, pluck="name") or [""]]},
		fields=["name", "application", "candidate_name", "status", "scheduled_datetime"],
		order_by="scheduled_datetime asc",
	)
	# Interview days, for the date filter; the sheet is usually one day's panel.
	counts = {}
	for iv in interviews:
		if iv.scheduled_datetime:
			day = iv.scheduled_datetime.date()
			counts[day] = counts.get(day, 0) + 1
	dates = [{"date": d, "count": n} for d, n in sorted(counts.items())]
	if interview_date:
		day = frappe.utils.getdate(interview_date)
		interviews = [iv for iv in interviews if iv.scheduled_datetime and iv.scheduled_datetime.date() == day]
	rubric = interview_rubric(interviews[0].name) if interviews else None
	if not rubric and job.track:
		name = frappe.db.get_value("Scoring Rubric Template", {"track": job.track, "stage": "Final Interview", "is_active": 1}, "name")
		rubric = frappe.get_doc("Scoring Rubric Template", name) if name else None
	criteria = [{"label": c.criterion_label, "max_score": flt(c.max_score), "guidance": c.guidance_text or ""} for c in (rubric.criteria if rubric else [])]
	max_total = sum(c["max_score"] for c in criteria)
	pass_percent = flt(rubric.pass_threshold_percent) if rubric else 50

	assessments = {}
	if interviews:
		for a in frappe.get_all(
			"Interview Assessment",
			filters={"interview": ["in", [iv.name for iv in interviews]]},
			fields=["name", "interview", "panelist", "total_score", "max_score", "verdict", "recommendation", "area_of_specialization", "entered_by"],
		):
			a.criteria = {
				r.criterion_label: r.score_given
				for r in frappe.get_all("Score Criterion Value", filters={"parent": a.name, "parenttype": "Interview Assessment"}, fields=["criterion_label", "score_given"])
			}
			assessments.setdefault(a.interview, {})[a.panelist] = a
	apps = {
		a.name: a
		for a in frappe.get_all(
			"Application",
			filters={"name": ["in", [iv.application for iv in interviews] or [""]]},
			fields=["name", "application_id", "status", "selection_outcome"],
		)
	}

	rows = []
	for iv in interviews:
		by_panellist = assessments.get(iv.name, {})
		scored = [by_panellist[m] for m in members if m in by_panellist]
		total_all = sum(flt(a.total_score) for a in scored)
		average = total_all / len(scored) if scored else 0
		verdicts = {v: sum(1 for a in scored if a.verdict == v) for v in VERDICTS}
		app = apps.get(iv.application) or frappe._dict()
		rows.append(
			{
				"interview": iv.name,
				"application": iv.application,
				"application_id": app.application_id,
				"candidate_name": iv.candidate_name,
				"interview_status": iv.status,
				"scheduled_datetime": iv.scheduled_datetime,
				"application_status": app.status,
				"outcome": app.selection_outcome or "",
				"scores": {
					m: {
						"total": flt(a.total_score), "criteria": a.criteria, "verdict": a.verdict, "comments": a.recommendation,
						"area": a.area_of_specialization, "entered_by": frappe.utils.get_fullname(a.entered_by) if a.entered_by else "",
					}
					for m, a in by_panellist.items() if m in members
				},
				"scored": len(scored),
				"total_of_all": round(total_all, 2),
				"average": round(average, 2),
				"percent": round(average * 100 / max_total, 1) if max_total else 0,
				"meets_bar": bool(scored) and max_total and average * 100 / max_total >= pass_percent,
				"verdicts": verdicts,
				"can_decide": app.status in OUTCOME_FROM,
			}
		)
	# Best average first; candidates nobody has scored yet at the end.
	rows.sort(key=lambda r: (r["scored"] == 0, -r["average"]))
	rank = 0
	for r in rows:
		if r["scored"]:
			rank += 1
			r["rank"] = rank
	interview_dates = sorted({iv.scheduled_datetime.date() for iv in interviews if iv.scheduled_datetime})
	if interview_date:
		heading_date = frappe.utils.getdate(interview_date)
	elif len(interview_dates) == 1:
		heading_date = interview_dates[0]
	else:
		heading_date = committee.scheduled_date if committee and committee.scheduled_date else (interview_dates[0] if interview_dates else None)
	return {
		"job": job,
		"interview_date": heading_date,
		"interview_dates": interview_dates,
		"dates": dates,
		"selected_date": str(frappe.utils.getdate(interview_date)) if interview_date else "",
		"criteria": criteria,
		"max_total": max_total,
		"pass_percent": pass_percent,
		"panel": [{"user": m, "full_name": frappe.utils.get_fullname(m)} for m in members],
		"rows": rows,
		"can_decide": 1 if _can_manage() else 0,
	}


@frappe.whitelist(methods=["POST"])
def set_selection_outcome(applications, outcome):
	"""Record the committee's decision for one or more candidates from the
	consolidated sheet: Selected / Waitlisted / Not Selected, or "" to clear
	it. Selected and Not Selected set the application status; a waitlisted
	candidate stays at Interview Completed (the portal shows the decision is
	being finalised). An open final interview is marked completed."""
	if not _can_manage():
		frappe.throw(_("Only the recruitment team can record the selection decision."), frappe.PermissionError)
	if outcome and outcome not in OUTCOMES:
		frappe.throw(_("Unknown outcome."))
	names = frappe.parse_json(applications) if isinstance(applications, str) else applications
	done, skipped = [], []
	for name in names or []:
		app = frappe.get_doc("Application", name)
		if app.status not in OUTCOME_FROM:
			skipped.append(f"{app.application_id}: {app.status}")
			continue
		if outcome:
			for iv in frappe.get_all("Interview", filters={"application": name, "round_type": "Final", "status": ["in", ["Scheduled", "Rescheduled"]]}, pluck="name"):
				frappe.db.set_value("Interview", iv, "status", "Completed")
		app.selection_outcome = outcome or None
		app.selection_decided_on = frappe.utils.now_datetime() if outcome else None
		app.status = {"Selected": "Selected", "Not Selected": "Not Selected"}.get(outcome, "Interview Completed")
		app.save(ignore_permissions=True)
		app.add_comment(
			"Info",
			_("Selection outcome: {0}, recorded by {1}.").format(_(outcome), frappe.utils.get_fullname(frappe.session.user))
			if outcome else _("Selection outcome cleared by {0}.").format(frappe.utils.get_fullname(frappe.session.user)),
		)
		done.append(name)
	return {"updated": len(done), "skipped": skipped}


@frappe.whitelist()
def export_consolidated(job_opening, interview_date=None):
	"""The consolidated score sheet as Excel, laid out like the workflow
	folder's sheet: a block of criteria + total per panellist, then the
	total of all and the average; a second sheet with every comment."""
	from io import BytesIO

	from openpyxl import Workbook
	from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
	from openpyxl.utils import get_column_letter

	data = get_consolidated(job_opening, interview_date)
	crit, panel, max_total = data["criteria"], data["panel"], data["max_total"]
	n_panel = len(panel)
	block = len(crit) + 1
	wb = Workbook()
	ws = wb.active
	ws.title = "Consolidated Score Sheet"
	thin = Side(style="thin", color="999999")
	border = Border(left=thin, right=thin, top=thin, bottom=thin)
	head_fill = PatternFill("solid", fgColor="F2E6E8")
	bold = Font(bold=True)
	center = Alignment(horizontal="center", vertical="center", wrap_text=True)

	tail = [f"Total of {n_panel}", f"Average for {max_total:g} Marks", "%", f"Meets {data['pass_percent']:g}%", "Recommendations", "Decision"]
	last_col = 2 + n_panel * block + len(tail)
	ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
	ws.cell(1, 1, "National Law School of India University, Bengaluru").font = Font(bold=True, size=14)
	ws.cell(1, 1).alignment = center
	days = data["interview_dates"]
	when = ", ".join(frappe.utils.formatdate(d, "dd.MM.yyyy") for d in days) if days else (frappe.utils.formatdate(data["interview_date"], "dd.MM.yyyy") if data["interview_date"] else "")
	ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
	ws.cell(2, 1, f"Position : {data['job']['job_title']}" + (f"          Date of Interview: {when}" if when else "")).font = bold
	ws.cell(2, 1).alignment = center

	for col, label in ((1, "Sl.\nNo."), (2, "Name of the Candidate")):
		ws.merge_cells(start_row=3, start_column=col, end_row=4, end_column=col)
		ws.cell(3, col, label)
	col = 3
	for p in panel:
		ws.merge_cells(start_row=3, start_column=col, end_row=3, end_column=col + block - 1)
		ws.cell(3, col, p["full_name"])
		for i, c in enumerate(crit):
			ws.cell(4, col + i, f"{c['label']}\n({c['max_score']:g} Marks)")
		ws.cell(4, col + len(crit), f"Total\n({max_total:g})")
		col += block
	for label in tail:
		ws.merge_cells(start_row=3, start_column=col, end_row=4, end_column=col)
		ws.cell(3, col, label)
		col += 1
	for r in (3, 4):
		for c in range(1, last_col + 1):
			cell = ws.cell(r, c)
			cell.font, cell.alignment, cell.fill, cell.border = bold, center, head_fill, border

	row_no = 5
	for i, r in enumerate(data["rows"], start=1):
		values = [i, f"{r['candidate_name']} ({r['application_id']})"]
		for p in panel:
			s = r["scores"].get(p["user"])
			values += [s["criteria"].get(c["label"]) if s else None for c in crit] + [s["total"] if s else None]
		v = r["verdicts"]
		values += [
			r["total_of_all"] if r["scored"] else None,
			r["average"] if r["scored"] else None,
			r["percent"] / 100 if r["scored"] else None,
			("Yes" if r["meets_bar"] else "No") if r["scored"] else "Not scored",
			", ".join(f"{n} {k}" for k, n in v.items() if n),
			r["outcome"],
		]
		for c, value in enumerate(values, start=1):
			cell = ws.cell(row_no, c, value)
			cell.border = border
			cell.alignment = Alignment(horizontal="left" if c == 2 else "center", vertical="center", wrap_text=True)
		ws.cell(row_no, last_col - 3).number_format = "0.0%"
		if r["outcome"] == "Selected":
			fill = PatternFill("solid", fgColor="DCFCE7")
		elif r["outcome"] == "Waitlisted":
			fill = PatternFill("solid", fgColor="FEF3C7")
		else:
			fill = None
		if fill:
			for c in range(1, last_col + 1):
				ws.cell(row_no, c).fill = fill
		row_no += 1

	ws.column_dimensions["A"].width = 6
	ws.column_dimensions["B"].width = 32
	for c in range(3, last_col + 1):
		ws.column_dimensions[get_column_letter(c)].width = 13
	ws.row_dimensions[4].height = 48
	ws.freeze_panes = "C5"

	notes = wb.create_sheet("Panel Comments")
	head = ["Candidate", "Candidate ID", "Panellist", "Total", "Recommendation", "Area of Specialization", "Additional Comments", "Entered By"]
	notes.append(head)
	for c in range(1, len(head) + 1):
		cell = notes.cell(1, c)
		cell.font, cell.fill, cell.border = bold, head_fill, border
	for r in data["rows"]:
		for p in panel:
			s = r["scores"].get(p["user"])
			if s:
				notes.append([r["candidate_name"], r["application_id"], p["full_name"], s["total"], s["verdict"] or "", s["area"] or "", s["comments"] or "", s["entered_by"]])
	for c, w in zip("ABCDEFGH", (28, 22, 24, 8, 18, 24, 60, 20)):
		notes.column_dimensions[c].width = w
	for row in notes.iter_rows(min_row=2):
		for cell in row:
			cell.alignment = Alignment(vertical="top", wrap_text=True)

	buffer = BytesIO()
	wb.save(buffer)
	suffix = f" - {frappe.utils.formatdate(interview_date, 'dd.MM.yyyy')}" if interview_date else ""
	frappe.local.response.filename = f"Consolidated Score Sheet - {data['job']['job_title']}{suffix}.xlsx"
	frappe.local.response.filecontent = buffer.getvalue()
	frappe.local.response.type = "download"


# ------------------------------------------------ what the panel looks at

def _doc(label, url):
	ext = (url or "").rsplit(".", 1)[-1].lower()
	kind = "pdf" if ext == "pdf" else "image" if ext in ("png", "jpg", "jpeg", "gif", "webp") else "other"
	return {"label": label, "file_url": url, "file_name": (url or "").rsplit("/", 1)[-1], "kind": kind}


def candidate_brief(application):
	"""What the workflow lists for the panel (step 16): CV, SOP, writing
	sample and other documents, qualifications, experience, the shortlisting
	note and how Round 1 went; not the whole application."""
	app = frappe.get_doc("Application", application)
	app.check_permission("read")
	docs = []
	for label, field in ((_("CV / Resume"), "resume_attachment"), (_("Statement of Purpose"), "sop_attachment"), (_("Additional attachment"), "additional_attachment")):
		if app.get(field):
			docs.append(_doc(label, app.get(field)))
	for row in app.get("documents") or []:
		if row.attachment and row.attachment not in [d["file_url"] for d in docs]:
			docs.append(_doc(row.document_type or _("Document"), row.attachment))
	institutions = {}
	for q in app.get("qualifications") or []:
		if q.institution and q.institution not in institutions:
			institutions[q.institution] = frappe.db.get_value("Institution Master", q.institution, "institution_name") or q.institution
	shortlisting = frappe.get_all("Shortlisting Score", filters={"application": application}, fields=["remarks", "total_score", "is_shortlisted"], order_by="modified desc", limit=1)
	round1 = frappe.get_all(
		"Interview", filters={"application": application, "round_type": "HR Interaction", "status": ["!=", "Cancelled"]},
		fields=["status", "scheduled_datetime"], order_by="scheduled_datetime desc", limit=1,
	)
	return {
		"full_name": frappe.db.get_value("Candidate", app.candidate, "full_name"),
		"application_id": app.application_id,
		"track": app.track,
		"overall_experience_years": app.get("overall_experience_years"),
		"relevant_experience_years": app.get("relevant_experience_years"),
		"qualifications": [
			{
				"level": q.degree_level, "degree": q.degree_name, "specialization": q.specialization, "year": q.year_of_graduation,
				"institution": institutions.get(q.institution) or q.other_institution or "", "score": q.percentage_or_cgpa or q.cgpa or "",
			}
			for q in app.get("qualifications") or []
		],
		"employment": [
			{"designation": e.designation, "employer": e.employer_name, "from_date": e.from_date, "to_date": e.to_date, "is_current": e.is_current}
			for e in app.get("employment_history") or []
		],
		"documents": docs,
		"shortlisting_note": (shortlisting[0].remarks if shortlisting else "") or "",
		"round1": {"status": round1[0].status, "date": round1[0].scheduled_datetime} if round1 else None,
	}


# --------------------------------------------- printable assessment forms

TRACK_FORM_NOTE = {
	"Admin": "This scoresheet will serve as an Interview Assessment Form for all non-teaching staff recruitment",
	"Faculty": "This scoresheet will serve as an Interview Assessment Form for all Teaching Staff recruitment, irrespective of levels",
	"Research": "This scoresheet will serve as an Interview Assessment Form for all research staff recruitment",
}

FORM_HTML = """
<style>
	body { font-family: Arial, sans-serif; font-size: 10.5px; color: #111; }
	.page { page-break-after: always; }
	.page:last-child { page-break-after: auto; }
	h1 { font-size: 15px; text-align: center; margin: 0 0 4px; }
	h2 { font-size: 12px; text-align: center; margin: 0 0 8px; letter-spacing: .5px; }
	ul { margin: 0 0 8px 16px; padding: 0; }
	table { width: 100%; border-collapse: collapse; }
	td, th { border: 1px solid #555; padding: 5px 6px; vertical-align: top; }
	th { background: #f2e6e8; text-align: center; font-weight: bold; }
	.q { font-weight: normal; font-size: 9px; color: #333; text-align: left; }
	.num { text-align: center; }
	.meta td { border: 1px solid #555; }
	.foot td { height: 26px; }
	.muted { color: #666; font-size: 9px; }
</style>
{% for f in forms %}
<div class="page">
	<h1>NATIONAL LAW SCHOOL OF INDIA UNIVERSITY, BENGALURU</h1>
	<h2>{{ position | upper }} | RECRUITMENT | INTERVIEW ASSESSMENT FORM</h2>
	<ul>
		<li>{{ note }}</li>
		<li>This scoresheet is to be completed by each member of the Selection Committee</li>
		<li>To be eligible for an offer, a candidate requires a minimum score of {{ pass_percent }}% of the total marks</li>
		<li>Whenever possible, please ensure that there is a waitlist recommended</li>
	</ul>
	<table class="meta" style="margin-bottom:8px">
		<tr><td style="width:50%"><b>PANELIST NAME:</b> {{ f.panellist }}</td><td><b>POSITION:</b> {{ position }}</td></tr>
		<tr><td></td><td><b>DATE:</b> {{ date }} &nbsp;&nbsp;&nbsp; <b>TIME:</b> {{ time }}</td></tr>
	</table>
	<table>
		<tr>
			<th style="width:4%">SL NO</th>
			<th style="width:14%">CANDIDATE NAME</th>
			{% if ask_area %}<th style="width:11%">AREA OF SPECIALIZATION</th>{% endif %}
			{% for c in criteria %}
			<th>{{ loop.index }}. {{ c.label | upper }}<br>(MAX SCORE = {{ c.max }})<div class="q">{{ c.guidance }}</div></th>
			{% endfor %}
			<th style="width:9%">TOTAL SCORE<br>({% for c in criteria %}{{ loop.index }}{% if not loop.last %} + {% endif %}{% endfor %})<br>(MAX SCORE = {{ max_total }})</th>
			<th style="width:18%">ADDITIONAL COMMENTS</th>
		</tr>
		{% for r in f.rows %}
		<tr>
			<td class="num">{{ loop.index }}</td>
			<td>{{ r.name }}<div class="muted">{{ r.id }}</div></td>
			{% if ask_area %}<td>{{ r.area }}</td>{% endif %}
			{% for v in r.scores %}<td class="num">{{ v }}</td>{% endfor %}
			<td class="num"><b>{{ r.total }}</b></td>
			<td>{{ r.comments }}</td>
		</tr>
		{% endfor %}
	</table>
	<table class="foot" style="margin-top:8px">
		<tr><td style="width:22%"><b>Recommendation of the Panelist:</b></td><td>{{ f.recommendation }}</td></tr>
		<tr><td><b>Signature of Panelist:</b></td><td>{{ f.signature }}</td></tr>
	</table>
</div>
{% endfor %}
"""


@frappe.whitelist()
def download_assessment_forms(job_opening, panelist=None, interview_date=None):
	"""The Interview Assessment Form as a PDF, one page per panellist with
	every candidate of the job, filled with the scores given so far (blank
	where not scored, for use on paper). Panellists get their own form; the
	recruitment team gets every panellist's, or one with `panelist`."""
	data = get_consolidated(job_opening, interview_date)
	members = [p for p in data["panel"]]
	if not _can_manage():
		members = [p for p in members if p["user"] == frappe.session.user]
	elif panelist:
		members = [p for p in members if p["user"] == panelist]
	if not members:
		frappe.throw(_("No panellist to print a form for."))
	job = data["job"]
	labels = {"Recommended": _("Recommended"), "Waitlist": _("Waitlist"), "Not Recommended": _("Not recommended")}
	forms = []
	for p in members:
		rows, by_verdict, confirmed = [], {}, []
		for r in data["rows"]:
			s = r["scores"].get(p["user"])
			rows.append(
				{
					"name": r["candidate_name"], "id": r["application_id"], "area": (s or {}).get("area") or "",
					"scores": [f"{flt(s['criteria'].get(c['label'])):g}" if s and s["criteria"].get(c["label"]) is not None else "" for c in data["criteria"]],
					"total": f"{s['total']:g}" if s else "", "comments": (s or {}).get("comments") or "",
				}
			)
			if s and s.get("verdict"):
				by_verdict.setdefault(s["verdict"], []).append(r["candidate_name"])
			if s:
				confirmed.append(s)
		recommendation = "; ".join(f"{labels[v]}: {', '.join(names)}" for v, names in by_verdict.items())
		entered = {s["entered_by"] for s in confirmed if s.get("entered_by")}
		signature = ""
		if confirmed:
			signature = _("Confirmed on the recruitment platform") + (f" ({_('entered by')} {', '.join(sorted(entered))})" if entered else "")
		forms.append({"panellist": p["full_name"], "rows": rows, "recommendation": recommendation, "signature": signature})

	times = sorted(r["scheduled_datetime"] for r in data["rows"] if r["scheduled_datetime"])
	when = times[0] if times else None
	html = frappe.render_template(
		FORM_HTML,
		{
			"forms": forms,
			"position": job.job_title,
			"note": TRACK_FORM_NOTE.get(job.track, TRACK_FORM_NOTE["Admin"]),
			"pass_percent": f"{data['pass_percent']:g}",
			"criteria": [{"label": c["label"], "max": f"{c['max_score']:g}", "guidance": c.get("guidance", "")} for c in data["criteria"]],
			"max_total": f"{data['max_total']:g}",
			"ask_area": job.track in SPECIALIZATION_TRACKS,
			"date": ", ".join(frappe.utils.formatdate(d, "dd MMM yyyy") for d in data["interview_dates"]) or (frappe.utils.formatdate(data["interview_date"], "dd MMM yyyy") if data["interview_date"] else ""),
			"time": frappe.utils.format_datetime(when, "h:mm a") if when else "",
		},
	)
	from frappe.utils.pdf import get_pdf

	pdf = get_pdf(html, {"orientation": "Landscape", "page-size": "A4", "margin-top": "10mm", "margin-bottom": "10mm", "margin-left": "8mm", "margin-right": "8mm"})
	who = members[0]["full_name"] if len(members) == 1 else _("All panellists")
	suffix = f" - {frappe.utils.formatdate(interview_date, 'dd.MM.yyyy')}" if interview_date else ""
	frappe.local.response.filename = f"Interview Assessment Form - {job.job_title} - {who}{suffix}.pdf"
	frappe.local.response.filecontent = pdf
	frappe.local.response.type = "download"
