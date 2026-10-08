# Copyright (c) 2026, NLSIU and contributors
# For license information, please see license.txt

"""Reports page: application charts and their Excel export.

Aggregates only (counts per group), behind the same gate as the
dashboard pipeline numbers. One function, _charts, builds every chart so
the page, the CSV and the Excel file always agree.

Each chart: {key, title, subtitle, kind, total, table, ...data}, where
kind is one of
  pie / donut   part-to-whole           slices: [{label, value, color, detail?}]
  bar           ranked, horizontal      slices
  funnel        stages reached          slices (shares are of the first stage)
  column        ordered bins            slices
  stacked       groups split by stage   categories, series: [{label, color, values}]
  line          applications over time  points: [{label, value}]
and table = {header, rows, percent: [column indexes holding fractions]}
is the same numbers as a grid (CSV, Excel)."""

import datetime
import io

import frappe
from frappe import _
from frappe.utils import add_days, add_months, date_diff, get_datetime, get_first_day, getdate, now_datetime, today

from pathways.api.dashboard import FUNNEL_STAGES, require_pipeline_access

# Charts are themed from the brand colour (Pathways Settings > Appearance,
# NLS maroon by default), in the same shades the app uses (theme.js).
DEFAULT_BRAND = "#920C24"
BRAND_SHADES = {300: 0.55, 400: 0.35, 500: 0.18, 600: 0.08, 700: 0, 900: -0.36}
# After the brand, the other categorical hues in a fixed order. With the
# maroon lead this order is validated colour-blind safe for neighbours,
# including a pie's last-to-first wrap.
OTHER_HUES = ["#2a78d6", "#eb6834", "#1baf7a", "#4a3aa7", "#eda100", "#008300", "#e87ba4"]
# Fills a slot only when a hue above is dropped for looking like the brand.
RESERVE_HUES = ["#e34948"]
NEUTRAL = "#898781"  # "Other" / "Not specified"
STATUS_GOOD, STATUS_WARNING, STATUS_CRITICAL = "#0ca30c", "#fab219", "#d03b3b"

# Application statuses grouped into eight stages (one colour each, the
# same in every chart).
STAGES = [
	("Screening", ["Submitted", "Under Review"]),
	("Shortlisted", ["Shortlisted"]),
	("Interview", ["Interview Scheduled", "Interview Completed"]),
	("Selected", ["Selected"]),
	("Offer", ["Offer Extended", "Offer Accepted", "Offer Declined"]),
	("Onboarding", ["Documents Pending", "Documents Verified"]),
	("Joined", ["Joined"]),
	("Closed", ["Not Selected", "Withdrawn"]),
]
STAGE_OF = {status: stage for stage, statuses in STAGES for status in statuses}
# No longer moving: left out of the ageing chart.
FINISHED = {"Joined", "Not Selected", "Withdrawn", "Offer Declined"}
AGE_BANDS = [(0, 7, "0–7 days"), (8, 14, "8–14 days"), (15, 30, "15–30 days"), (31, 60, "31–60 days"), (61, None, "Over 60 days")]


def _hex(rgb):
	return "#{:02x}{:02x}{:02x}".format(*rgb)


def _brand_scale(color):
	"""Shades of the brand colour: 700 is the colour itself, lighter ones
	blend toward white, darker toward black (as theme.js brandScale)."""
	value = (color or "").strip().lstrip("#")
	if len(value) != 6 or any(c not in "0123456789abcdefABCDEF" for c in value):
		value = DEFAULT_BRAND.lstrip("#")
	base = [int(value[i : i + 2], 16) for i in (0, 2, 4)]
	scale = {}
	for shade, amount in BRAND_SHADES.items():
		target = 255 if amount >= 0 else 0
		scale[shade] = _hex([round(c + (target - c) * abs(amount)) for c in base])
	return scale


def _oklab(color):
	def linear(c):
		c /= 255
		return c / 12.92 if c <= 0.04045 else ((c + 0.055) / 1.055) ** 2.4

	r, g, b = (linear(int(color[i : i + 2], 16)) for i in (1, 3, 5))
	l = (0.4122214708 * r + 0.5363325363 * g + 0.0514459929 * b) ** (1 / 3)
	m = (0.2119034982 * r + 0.6806995451 * g + 0.1073969566 * b) ** (1 / 3)
	s = (0.0883024619 * r + 0.2817188376 * g + 0.6299787005 * b) ** (1 / 3)
	return (
		0.2104542553 * l + 0.7936177850 * m - 0.0040720468 * s,
		1.9779984951 * l - 2.4285922050 * m + 0.4505937099 * s,
		0.0259040371 * l + 0.7827717662 * m - 0.8086757660 * s,
	)


def _distance(a, b):
	"""Perceptual colour difference (OKLab x 100); under 15 reads as the same colour."""
	return 100 * sum((x - y) ** 2 for x, y in zip(_oklab(a), _oklab(b))) ** 0.5


def _theme():
	"""{primary, ramp, palette} for the configured brand colour.

	palette[0] is the brand (shade 600: shade 700 is too dark to sit among
	the other hues); a hue that looks too like it is dropped, so no two
	slots look alike if the brand is changed from maroon."""
	scale = _brand_scale(frappe.get_cached_doc("Pathways Settings").get("brand_color"))
	lead = scale[600]
	return frappe._dict(
		primary=scale[700],
		ramp=[scale[s] for s in (300, 400, 500, 700, 900)],
		palette=[lead] + [c for c in OTHER_HUES + RESERVE_HUES if _distance(c, lead) >= 15][: len(OTHER_HUES)],
	)


def _filters(track=None, job_opening=None, from_date=None, to_date=None):
	"""Application filters; every given filter applies (AND)."""
	filters = {}
	jobs = None
	if track:
		jobs = set(frappe.get_all("Job Opening", filters={"track": track}, pluck="name"))
	if job_opening:
		jobs = {job_opening} if jobs is None else jobs & {job_opening}
	if jobs is not None:
		# An empty list would mean "no filter" to frappe: match nothing instead.
		filters["job_opening"] = ["in", list(jobs) or [""]]

	start = getdate(from_date) if from_date else None
	end = getdate(to_date) if to_date else None
	if start and end and start > end:
		frappe.throw(_("The From date must be on or before the To date."))
	if start and end:
		filters["application_date"] = ["between", [start, end]]
	elif start:
		filters["application_date"] = [">=", start]
	elif end:
		filters["application_date"] = ["<=", end]
	return filters


def _job_filters(track=None, job_opening=None):
	f = {}
	if track:
		f["track"] = track
	if job_opening:
		f["name"] = job_opening
	return f


def _slices(counts, palette, order=None):
	"""[{label, value, color}] for a pie.

	With a fixed `order` of at most eight entities, each keeps its slot
	colour whatever the filters (colour follows the entity). Otherwise the
	largest groups take the slots in rank order and the rest fold into
	"Other (n)"."""
	counts = {k: v for k, v in counts.items() if v}
	unspecified = counts.pop(None, 0)
	if order is not None and len(order) <= len(palette):
		slices = [{"label": key, "value": counts.pop(key, 0), "color": palette[i]} for i, key in enumerate(order)]
		# Values outside the known list (e.g. a renamed master).
		slices += [{"label": k, "value": v, "color": NEUTRAL} for k, v in sorted(counts.items(), key=lambda kv: -kv[1])]
	else:
		ranked = sorted(counts.items(), key=lambda kv: (-kv[1], str(kv[0])))
		keep = ranked if len(ranked) <= len(palette) else ranked[: len(palette) - 1]
		rest = ranked[len(keep) :]
		slices = [{"label": key, "value": value, "color": palette[i]} for i, (key, value) in enumerate(keep)]
		if rest:
			slices.append(
				{
					"label": _("Other ({0})").format(len(rest)),
					"value": sum(v for _k, v in rest),
					"color": NEUTRAL,
					"detail": [{"label": k, "value": v} for k, v in rest],
				}
			)
	if unspecified:
		slices.append({"label": _("Not specified"), "value": unspecified, "color": NEUTRAL})
	return slices


def _count(rows, key):
	counts = {}
	for row in rows:
		value = key(row) or None
		counts[value] = counts.get(value, 0) + 1
	return counts


def _select_options(doctype, fieldname):
	return [o for o in (frappe.get_meta(doctype).get_field(fieldname).options or "").split("\n") if o]


def _slice_chart(key, kind, title, subtitle, slices, first_column):
	"""A chart drawn from slices, with its table."""
	total = sum(s["value"] for s in slices)
	if kind == "funnel":
		# Shares are of the first stage (every application), not of the sum.
		total = slices[0]["value"] if slices else 0
		header = [first_column, _("Applications"), _("% of Applied")]
	else:
		header = [first_column, _("Applications"), _("Share")]
	rows = [[s["label"], s["value"], (s["value"] / total) if total else 0] for s in slices]
	return {
		"key": key,
		"kind": kind,
		"title": title,
		"subtitle": subtitle,
		"slices": slices,
		"total": total,
		"table": {"header": header, "rows": rows, "percent": [2]},
	}


def _timeline(apps, from_date, to_date):
	"""Applications per day, week (from Monday) or month, gaps filled with 0."""
	dates = [getdate(a.application_date or get_datetime(a.creation)) for a in apps]
	if not dates:
		return [], "day"
	start = getdate(from_date) if from_date else min(dates)
	end = getdate(to_date) if to_date else max(dates)
	span = date_diff(end, start)
	if span <= 45:
		unit, bucket = "day", lambda d: d
	elif span <= 7 * 30:
		unit, bucket = "week", lambda d: d - datetime.timedelta(days=d.weekday())
	else:
		unit, bucket = "month", lambda d: get_first_day(d)

	counts = {}
	for d in dates:
		counts[bucket(d)] = counts.get(bucket(d), 0) + 1
	points = []
	current, last = bucket(start), bucket(end)
	while current <= last:
		label = current.strftime("%b %Y") if unit == "month" else current.strftime("%d %b %Y")
		points.append({"label": label, "date": current.isoformat(), "value": counts.get(current, 0)})
		current = add_months(current, 1) if unit == "month" else add_days(current, 7 if unit == "week" else 1)
	return points, unit


def _charts(filters, job_filters, from_date=None, to_date=None):
	theme = _theme()
	# One colour per stage, the same in every chart.
	stage_color = {stage: theme.palette[i] if i < len(theme.palette) else NEUTRAL for i, (stage, _s) in enumerate(STAGES)}
	apps = frappe.get_all(
		"Application",
		filters=filters,
		fields=["name", "status", "eligibility_status", "job_opening", "source", "candidate", "application_date", "creation"],
		limit_page_length=0,
	)
	jobs = {
		j.name: j
		for j in frappe.get_all("Job Opening", filters=job_filters, fields=["name", "job_title", "track"], order_by="creation desc")
	}
	missing = {a.job_opening for a in apps} - set(jobs)
	if missing:
		for j in frappe.get_all("Job Opening", filters={"name": ["in", list(missing)]}, fields=["name", "job_title", "track"]):
			jobs[j.name] = j
	candidates = {}
	if apps:
		candidates = {
			c.name: c
			for c in frappe.get_all(
				"Candidate", filters={"name": ["in", list({a.candidate for a in apps})]}, fields=["name", "gender", "category"]
			)
		}
	track_of = lambda a: jobs.get(a.job_opening, {}).get("track")  # noqa: E731
	charts = []

	# 1. Applications over time
	points, unit = _timeline(apps, from_date, to_date)
	period = {"day": _("Date"), "week": _("Week Starting"), "month": _("Month")}[unit]
	charts.append(
		{
			"key": "timeline",
			"kind": "line",
			"color": theme.primary,
			"title": _("Applications Over Time"),
			"subtitle": _("Applications received per {0}.").format(_(unit)),
			"points": points,
			"unit": unit,
			"total": sum(p["value"] for p in points),
			"table": {"header": [period, _("Applications")], "rows": [[p["label"], p["value"]] for p in points], "percent": []},
		}
	)

	# 2. Funnel: applications that REACHED each stage (dashboard's definition).
	by_status = _count(apps, lambda a: a.status)
	funnel = [
		{"label": _(label), "value": sum(by_status.get(s, 0) for s in statuses), "color": theme.ramp[min(i, len(theme.ramp) - 1)]}
		for i, (label, statuses) in enumerate(FUNNEL_STAGES)
	]
	charts.append(
		_slice_chart(
			"funnel", "funnel", _("Recruitment Funnel"), _("Applications that reached each stage, as a share of all applied."),
			funnel, _("Stage"),
		)
	)

	# 3. Current status, grouped into stages; the tooltip lists the statuses inside.
	stage_slices = []
	for stage, statuses in STAGES:
		detail = [{"label": s, "value": by_status[s]} for s in statuses if by_status.get(s)]
		stage_slices.append(
			{
				"label": _(stage),
				"value": sum(d["value"] for d in detail),
				"color": stage_color[stage],
				"detail": detail if len(statuses) > 1 else [],
			}
		)
	other = {s: n for s, n in by_status.items() if s not in STAGE_OF}
	if other:
		stage_slices.append(
			{
				"label": _("Other"),
				"value": sum(other.values()),
				"color": NEUTRAL,
				"detail": [{"label": s or _("Not specified"), "value": n} for s, n in other.items()],
			}
		)
	charts.append(
		_slice_chart("status", "donut", _("Application Status"), _("Where applications stand right now."), stage_slices, _("Stage"))
	)

	# 4. Each track split by stage.
	tracks = frappe.get_all("Recruitment Track", pluck="name", order_by="creation asc")
	per_track = {}
	for a in apps:
		counts = per_track.setdefault(track_of(a) or _("Not specified"), {})
		stage = STAGE_OF.get(a.status, "Other")
		counts[stage] = counts.get(stage, 0) + 1
	categories = [t for t in tracks if t in per_track] + sorted(t for t in per_track if t not in tracks)
	stage_names = [s for s, _x in STAGES] + (["Other"] if other else [])
	series = [
		{
			"label": _(stage),
			"color": stage_color.get(stage, NEUTRAL),
			"values": [per_track[c].get(stage, 0) for c in categories],
		}
		for stage in stage_names
	]
	series = [s for s in series if any(s["values"])]  # colours stay fixed per stage
	charts.append(
		{
			"key": "track_stage",
			"kind": "stacked",
			"title": _("Track × Stage"),
			"subtitle": _("Applications in each recruitment track, by current stage."),
			"categories": categories,
			"series": series,
			"total": len(apps),
			"table": {
				"header": [_("Track")] + [s["label"] for s in series] + [_("Total")],
				"rows": [[c] + [s["values"][i] for s in series] + [sum(s["values"][i] for s in series)] for i, c in enumerate(categories)],
				"percent": [],
			},
		}
	)

	# 5. Days in the current status, for applications still moving.
	active = [a for a in apps if a.status not in FINISHED]
	last_change = {}
	if active:
		for row in frappe.get_all(
			"Application Status History",
			filters={"application": ["in", [a.name for a in active]]},
			fields=["application", {"MAX": "changed_on", "as": "changed_on"}],
			group_by="application",
		):
			last_change[row.application] = row.changed_on
	bands = {label: 0 for _lo, _hi, label in AGE_BANDS}
	now = today()
	for a in active:
		since = last_change.get(a.name) or a.application_date or a.creation
		days = max(date_diff(now, getdate(get_datetime(since))), 0)
		for lo, hi, label in AGE_BANDS:
			if days >= lo and (hi is None or days <= hi):
				bands[label] += 1
				break
	charts.append(
		_slice_chart(
			"aging", "column", _("Application Ageing"), _("Days open applications have spent in their current status."),
			[{"label": _(label), "value": bands[label], "color": theme.ramp[i]} for i, (_lo, _hi, label) in enumerate(AGE_BANDS)],
			_("Days in Status"),
		)
	)

	# 6. Source, ranked: too many sources for a pie.
	sources = _count(apps, lambda a: a.source)
	unspecified = sources.pop(None, 0)
	source_slices = [
		{"label": k, "value": v, "color": theme.primary} for k, v in sorted(sources.items(), key=lambda kv: (-kv[1], kv[0]))
	]
	if unspecified:
		source_slices.append({"label": _("Not specified"), "value": unspecified, "color": NEUTRAL})
	charts.append(
		_slice_chart("source", "bar", _("Applications by Source"), _("Where applicants heard about the opening, most first."), source_slices, _("Source"))
	)

	# 7. Category in its own (reservation) order; missing answers in grey.
	categories_count = _count(apps, lambda a: candidates.get(a.candidate, {}).get("category"))
	category_slices = [
		{"label": c, "value": categories_count.pop(c, 0), "color": theme.primary} for c in _select_options("Candidate", "category")
	]
	unknown = categories_count.pop(None, 0)
	category_slices += [{"label": k, "value": v, "color": theme.primary} for k, v in categories_count.items()]
	if unknown:
		category_slices.append({"label": _("Not specified"), "value": unknown, "color": NEUTRAL})
	charts.append(
		_slice_chart("category", "column", _("Applicants by Category"), _("Category declared on each application."), category_slices, _("Category"))
	)

	# 8. Eligibility (status colours, always shown with their labels).
	eligibility = _count(apps, lambda a: a.eligibility_status if a.eligibility_status in ("Eligible", "Not Eligible") else "To check")
	charts.append(
		_slice_chart(
			"eligibility", "donut", _("Eligibility Check"), _("Applications checked against each job's eligibility criteria."),
			[
				{"label": _("Eligible"), "value": eligibility.get("Eligible", 0), "color": STATUS_GOOD},
				{"label": _("Not eligible"), "value": eligibility.get("Not Eligible", 0), "color": STATUS_CRITICAL},
				{"label": _("To check"), "value": eligibility.get("To check", 0), "color": STATUS_WARNING},
			],
			_("Eligibility"),
		)
	)

	# 9. Gender: a small part-to-whole, so a pie.
	charts.append(
		_slice_chart(
			"gender", "pie", _("Applicants by Gender"), _("Gender declared on each application."),
			_slices(
				_count(apps, lambda a: candidates.get(a.candidate, {}).get("gender")),
				theme.palette,
				order=_select_options("Candidate", "gender"),
			),
			_("Gender"),
		)
	)

	# 10. Every job in the selection (zero included), most applications first.
	per_job = _count(apps, lambda a: a.job_opening)
	job_slices = sorted(
		(
			{"label": job.job_title or name, "value": per_job.get(name, 0), "color": theme.primary, "id": name, "track": job.track}
			for name, job in jobs.items()
		),
		key=lambda s: (-s["value"], s["label"].lower()),
	)
	# 10. Positions: every job opening's applications under its position.
	pos_of = dict(frappe.get_all("Job Opening", filters={"name": ["in", list(jobs) or [""]]}, fields=["name", "position"], as_list=True))
	pos_titles = dict(frappe.get_all("Position", fields=["name", "position_title"], as_list=True))
	per_position = {}
	for s_ in job_slices:
		pos = pos_of.get(s_["id"]) or _("No position")
		per_position[pos] = per_position.get(pos, 0) + s_["value"]
	position_slices = sorted(
		(
			{"label": f"{pos} · {pos_titles[pos]}" if pos in pos_titles else pos, "value": n, "color": theme.primary, "id": pos}
			for pos, n in per_position.items()
		),
		key=lambda s_: (-s_["value"], s_["label"].lower()),
	)
	charts.append(
		_slice_chart(
			"position", "bar", _("Applications by Position"), _("Every position in this selection, most applications first."),
			position_slices, _("Position"),
		)
	)

	# Status, Track x Stage, Ageing and Gender repeat the Dashboard (cards,
	# Recruitment Pipeline, Needs Attention) or are too small to chart.
	keep = ("timeline", "funnel", "eligibility", "position", "source", "category")
	charts = sorted((c for c in charts if c["key"] in keep), key=lambda c: keep.index(c["key"]))
	return {"total": len(apps), "charts": charts}


LIST_FILTERS = (
	("positions", "Position"),
	("tracks", "Track"),
	("departments", "Department"),
	("employment_types", "Employment Type"),
	("job_openings", "Job Opening"),
	("statuses", "Application Status"),
	("eligibility", "Eligibility"),
)


def _resolve(filters=None, track=None, job_opening=None, from_date=None, to_date=None):
	"""(application filters, job filters, from, to, summary) for either the
	filter bar (`filters`, as on the Dashboard) or the older single filters."""
	from pathways.api.dashboard import build_filters

	f = frappe.parse_json(filters) if isinstance(filters, str) else (filters or {})
	if not f:
		f = {
			"tracks": [track] if track else [],
			"job_openings": [job_opening] if job_opening else [],
			"from_date": from_date,
			"to_date": to_date,
		}
	start, end = f.get("from_date") or None, f.get("to_date") or None
	if start and end and getdate(start) > getdate(end):
		frappe.throw(_("The From date must be on or before the To date."))
	app_filters, job_filters = build_filters({**f, "from_date": None, "to_date": None})
	if start and end:
		app_filters["application_date"] = ["between", [getdate(start), getdate(end)]]
	elif start:
		app_filters["application_date"] = [">=", getdate(start)]
	elif end:
		app_filters["application_date"] = ["<=", getdate(end)]

	titles = dict(frappe.get_all("Job Opening", fields=["name", "job_title"], as_list=True))
	rows = []
	for key, label in LIST_FILTERS:
		values = f.get(key) or []
		if key == "job_openings":
			values = [titles.get(v, v) for v in values]
		rows.append((_(label), ", ".join(values) if values else _("All")))
	rows.append((_("Applied From"), getdate(start) if start else _("Any date")))
	rows.append((_("Applied To"), getdate(end) if end else _("Any date")))
	return app_filters, job_filters, start, end, rows


def _options():
	# Positions that have a job opening: the others have nothing to report.
	used = set(frappe.get_all("Job Opening", filters={"position": ["is", "set"]}, pluck="position"))
	positions = [p for p in frappe.get_all("Position", fields=["name", "position_title"], order_by="name asc") if p.name in used]
	return {
		"positions": [{"value": p.name, "label": f"{p.name} · {p.position_title}"} for p in positions],
		"tracks": frappe.get_all("Recruitment Track", pluck="name", order_by="name asc"),
		"departments": frappe.get_all("Department", filters={"is_active": 1}, pluck="name", order_by="name asc"),
		"employment_types": [o for o in (frappe.get_meta("Job Opening").get_field("employment_type").options or "").split("\n") if o],
		"statuses": [o for o in (frappe.get_meta("Application").get_field("status").options or "").split("\n") if o],
		"eligibility": ["Pending", "Eligible", "Not Eligible"],
		"jobs": [
			{"value": j.name, "label": j.job_title or j.name, "track": j.track, "department": j.department, "position": j.position,
				"employment_type": j.employment_type}
			for j in frappe.get_all("Job Opening", fields=["name", "job_title", "track", "department", "position", "employment_type"], order_by="job_title asc")
		],
	}


@frappe.whitelist()
def get_report_charts(filters=None, track=None, job_opening=None, from_date=None, to_date=None):
	"""Every Reports page chart for the filters, plus the filter options."""
	require_pipeline_access()
	app_filters, job_filters, start, end, _rows = _resolve(filters, track, job_opening, from_date, to_date)
	data = _charts(app_filters, job_filters, start, end)
	data["options"] = _options()
	return data


@frappe.whitelist()
def export_report(filters=None, scope="all", positions=None, layout="single", include_charts=1, track=None, job_opening=None, from_date=None, to_date=None):
	"""Excel export from the Reports page.

	scope: "all" (everything the filters match) or "positions" (only the
	chosen positions). layout: "single" (all applications on one sheet) or
	"position_wise" (one sheet per position). include_charts adds the
	Summary, one sheet per chart and the pipeline table first."""
	from pathways.utils.application_export import add_application_sheets

	require_pipeline_access()
	f = frappe.parse_json(filters) if isinstance(filters, str) else dict(filters or {})
	if scope == "positions":
		positions = frappe.parse_json(positions) if isinstance(positions, str) else (positions or [])
		if not positions:
			frappe.throw(_("Choose at least one position to export."))
		f["positions"] = positions
	if layout not in ("single", "position_wise"):
		frappe.throw(_("Unknown export layout."))

	app_filters, job_filters, start, end, rows = _resolve(f, track, job_opening, from_date, to_date)
	names = frappe.get_list("Application", filters=app_filters, pluck="name", order_by="creation desc", limit_page_length=0)
	if not names:
		frappe.throw(_("No applications match these filters."))

	rows.append((_("Layout"), _("One sheet per position") if layout == "position_wise" else _("All applications on one sheet")))
	jobs = set(frappe.get_all("Job Opening", filters=job_filters, pluck="name")) if job_filters else None
	if frappe.utils.cint(include_charts):
		data = _charts(app_filters, job_filters, start, end)
		wb = _workbook(data, {"rows": rows, "jobs": jobs})
	else:
		from openpyxl import Workbook

		wb = Workbook()
		wb.active.title = _("Summary")
		wb.active.append([_("Applications Export")])
		for label, value in rows:
			wb.active.append([label, value])
	used = {name.lower() for name in wb.sheetnames}
	# Rows sorted by position, then stage, then eligibility, so each count on
	# the Summary links to one block of rows.
	info = _application_info(names)
	names = sorted(names, key=lambda n: info[n]["sort"])
	extra = {n: {_("Position"): info[n]["position_label"], _("Stage"): info[n]["stage"]} for n in names}
	locations = {}
	sheets, total = add_application_sheets(wb, names, layout, used, extra=extra, locations=locations)

	# List the data sheets on the Summary too.
	from openpyxl.styles import Font
	from openpyxl.worksheet.hyperlink import Hyperlink

	summary = wb[_("Summary")]
	summary.append([])
	summary.append([_("Application data"), _("{0} applications").format(total)])
	summary.cell(row=summary.max_row, column=1).font = Font(bold=True)
	for sheet_name, _ref, title, count in sheets:
		summary.append([sheet_name, f"{title} ({count})"])
		link = summary.cell(row=summary.max_row, column=1)
		link.hyperlink = Hyperlink(ref=link.coordinate, location=f"'{sheet_name.replace(chr(39), chr(39) * 2)}'!A1")
		link.font = Font(color="0563C1", underline="single")

	_colour_stage_cells(wb, [sheet for sheet, *_rest in sheets])
	_position_stage_table(wb, summary, names, info, locations)

	buffer = io.BytesIO()
	wb.save(buffer)
	suffix = "by-position" if layout == "position_wise" else "all"
	frappe.local.response.filename = f"recruitment-report-{suffix}-{today()}.xlsx"
	frappe.local.response.filecontent = buffer.getvalue()
	frappe.local.response.type = "download"


# Stage columns on the Summary's position table (Closed split in two).
SUMMARY_STAGES = [s for s in STAGES if s[0] != "Closed"] + [("Not Selected", ["Not Selected"]), ("Withdrawn", ["Withdrawn"])]


def _application_info(names):
	"""{application: position, stage, eligibility and a sort key}."""
	apps = frappe.get_all(
		"Application",
		filters={"name": ["in", names or [""]]},
		fields=["name", "job_opening", "status", "eligibility_status", "creation"],
		limit_page_length=0,
	)
	jobs = dict(
		frappe.get_all("Job Opening", filters={"name": ["in", list({a.job_opening for a in apps}) or [""]]}, fields=["name", "position"], as_list=True)
	)
	titles = dict(frappe.get_all("Position", fields=["name", "position_title"], as_list=True))
	counts = {}
	for a in apps:
		counts[jobs.get(a.job_opening) or ""] = counts.get(jobs.get(a.job_opening) or "", 0) + 1
	stage_index = {stage: i for i, (stage, _x) in enumerate(SUMMARY_STAGES)}
	elig_index = {"Eligible": 0, "Not Eligible": 1}
	out = {}
	for a in apps:
		pos = jobs.get(a.job_opening) or ""
		stage = next((st for st, statuses in SUMMARY_STAGES if a.status in statuses), _("Other"))
		elig = a.eligibility_status if a.eligibility_status in ("Eligible", "Not Eligible") else "Pending"
		out[a.name] = {
			"position": pos or _("No position"),
			"position_label": f"{pos} · {titles[pos]}" if pos in titles else (pos or _("No position")),
			"title": titles.get(pos, ""),
			"stage": stage,
			"eligibility": elig,
			# Same order as the Summary table: most applications first.
			"sort": (-counts[pos], pos, stage_index.get(stage, 99), elig_index.get(elig, 2), str(a.creation)),
		}
	return out


# Light fills per stage, on the Stage column and the Details headings.
STAGE_FILLS = {
	"Screening": "FFF4E5", "Shortlisted": "E6F4EA", "Interview": "E8F0FE", "Selected": "D9F2E3", "Offer": "D9F2E3",
	"Onboarding": "D9F2E3", "Joined": "C8EBD5", "Not Selected": "FDECEA", "Withdrawn": "EEEEEE",
	"Eligible": "E6F4EA", "Not Eligible": "FDECEA", "Pending": "FFF4E5",
}


def _colour_stage_cells(wb, sheet_names):
	from openpyxl.styles import PatternFill

	for name in sheet_names:
		ws = wb[name]
		header = [c.value for c in ws[1]]
		if "Stage" not in header:
			continue
		col = header.index("Stage") + 1
		for row in range(2, ws.max_row + 1):
			cell = ws.cell(row=row, column=col)
			colour = STAGE_FILLS.get(cell.value)
			if colour:
				cell.fill = PatternFill("solid", fgColor=colour)


DETAIL_COLUMNS = [
	("application_id", "Application ID"), ("candidate_name", "Candidate"), ("email", "Email"), ("mobile", "Mobile"),
	("job_title", "Job Opening"), ("stage", "Stage"), ("status", "Status"), ("eligibility", "Eligibility"),
	("reason", "Reason Not Eligible"), ("applied_on", "Applied On"),
]


def _position_stage_table(wb, ws, names, info, locations):
	"""Summary sheet: for each position, how many applied, their
	eligibility, and how many are at each stage, with a Total row. Each
	count links to its own section on the Details sheet, which lists just
	those applications."""
	from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
	from openpyxl.utils import get_column_letter
	from openpyxl.worksheet.hyperlink import Hyperlink

	stage_keys = [stage for stage, _x in SUMMARY_STAGES]
	keys = ["Eligible", "Not Eligible", "Pending"] + stage_keys
	key_label = {"Pending": _("To Check")}
	rows = {}
	for n in names:
		i = info[n]
		r = rows.setdefault(i["position"], {"title": i["title"], "apps": [], "jobs": set(), **{k: [] for k in keys}})
		r["apps"].append(n)
		r["jobs"].add(frappe.db.get_value("Application", n, "job_opening"))
		r[i["eligibility"]].append(n)
		if i["stage"] in r:
			r[i["stage"]].append(n)
	order = sorted(rows, key=lambda p: (-len(rows[p]["apps"]), p))

	header = [_("Position"), _("Position Title"), _("Job Openings"), _("Applications"), _("Eligible"), _("Not Eligible"), _("To Check")]
	header += [_(stage) for stage in stage_keys]

	ws.append([])
	ws.append([_("Applications by position and stage")])
	ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=12)
	ws.append([_("Click a number to see just those applications (Details sheet).")])
	ws.cell(row=ws.max_row, column=1).font = Font(italic=True, color="52514E")
	ws.append(header)
	head_row = ws.max_row
	brand = _theme().primary[1:].upper()
	fill = PatternFill("solid", fgColor=brand)
	for c in range(1, len(header) + 1):
		cell = ws.cell(row=head_row, column=c)
		cell.font = Font(bold=True, color="FFFFFF")
		cell.fill = fill
		cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center" if c > 2 else "left")
	ws.row_dimensions[head_row].height = 32

	# Rows of the table: (summary row, [(column, heading, apps)])
	links = []
	for pos in order:
		r = rows[pos]
		ws.append([pos, r["title"], len(r["jobs"]), len(r["apps"]), *[len(r[k]) for k in keys]])
		row = ws.max_row
		for c in range(1, 3):
			ws.cell(row=row, column=c).data_type = "s"
		name = f"{pos} · {r['title']}" if r["title"] else pos
		cells = [(4, f"{name} — {_('All applications')}", r["apps"])]
		cells += [(c, f"{name} — {_(key_label.get(k, k))}", r[k]) for c, k in enumerate(keys, start=5) if r[k]]
		links.append((row, cells))

	ws.append([_("Total"), "", sum(len(r["jobs"]) for r in rows.values()), sum(len(r["apps"]) for r in rows.values()),
		*[sum(len(r[k]) for r in rows.values()) for k in keys]])
	total_row = ws.max_row
	line = Side(style="thin", color="999999")
	for c in range(1, len(header) + 1):
		cell = ws.cell(row=total_row, column=c)
		cell.font = Font(bold=True)
		cell.border = Border(top=line)
	all_apps = [n for p in order for n in rows[p]["apps"]]
	total_cells = [(4, _("All positions — All applications"), all_apps)]
	for c, k in enumerate(keys, start=5):
		apps = [n for p in order for n in rows[p][k]]
		if apps:
			total_cells.append((c, f"{_('All positions')} — {_(key_label.get(k, k))}", apps))
	links.append((total_row, total_cells))
	for c in range(3, len(header) + 1):
		ws.column_dimensions[get_column_letter(c)].width = max(ws.column_dimensions[get_column_letter(c)].width or 0, 12)

	# ---- Details sheet: one highlighted section per count
	details = _details_rows(names)
	dws = wb.create_sheet(_("Details"), index=wb.sheetnames.index(ws.title) + 1)
	dws["A1"] = _("Applications behind each number on the Summary")
	dws["A1"].font = Font(bold=True, size=13)
	link_font = Font(color="0563C1", underline="single")
	width = len(DETAIL_COLUMNS)
	last_col = get_column_letter(width)
	summary_ref = ws.title.replace("'", "''")
	head_fill = PatternFill("solid", fgColor="F3F4F6")
	row = 3
	for summary_row, cells in links:
		for col, heading, apps in cells:
			start = row
			title = dws.cell(row=row, column=1, value=f"{heading} ({len(apps)})")
			title.font = Font(bold=True, color="FFFFFF")
			for c in range(1, width + 1):
				dws.cell(row=row, column=c).fill = fill
			back = dws.cell(row=row, column=width, value=_("↑ Back to Summary"))
			back.hyperlink = Hyperlink(ref=back.coordinate, location=f"'{summary_ref}'!{get_column_letter(col)}{summary_row}")
			back.font = Font(color="FFFFFF", underline="single")
			back.alignment = Alignment(horizontal="right")
			row += 1
			for c, (_key, label) in enumerate(DETAIL_COLUMNS, start=1):
				cell = dws.cell(row=row, column=c, value=_(label))
				cell.font = Font(bold=True)
				cell.fill = head_fill
			row += 1
			for n in apps:
				d = details.get(n, {})
				for c, (key, _label) in enumerate(DETAIL_COLUMNS, start=1):
					cell = dws.cell(row=row, column=c, value=d.get(key))
					if isinstance(cell.value, str):
						cell.data_type = "s"
					if key == "stage" and STAGE_FILLS.get(d.get("stage")):
						cell.fill = PatternFill("solid", fgColor=STAGE_FILLS[d["stage"]])
					if key == "eligibility" and STAGE_FILLS.get(d.get("eligibility_key")):
						cell.fill = PatternFill("solid", fgColor=STAGE_FILLS[d["eligibility_key"]])
					if key == "applied_on" and cell.value:
						cell.number_format = "DD-MMM-YYYY"
				row += 1
			end = row - 1
			# The Summary number opens this section, selected.
			target = ws.cell(row=summary_row, column=col)
			target.hyperlink = Hyperlink(ref=target.coordinate, location=f"'{dws.title}'!A{start}:{last_col}{end}")
			target.font = Font(bold=target.font.bold, color="0563C1", underline="single")
			row += 1  # gap
	for c, w in enumerate((20, 24, 30, 16, 34, 14, 18, 14, 30, 14), start=1):
		dws.column_dimensions[get_column_letter(c)].width = w
	dws.freeze_panes = "A2"


def _details_rows(names):
	apps = frappe.get_all(
		"Application",
		filters={"name": ["in", names or [""]]},
		fields=[
			"name", "application_id", "candidate.full_name as candidate_name", "candidate.email as email",
			"candidate.mobile_number as mobile", "job_opening.job_title as job_title", "status", "eligibility_status",
			"eligibility_reason", "application_date",
		],
		limit_page_length=0,
	)
	out = {}
	for a in apps:
		elig = a.eligibility_status if a.eligibility_status in ("Eligible", "Not Eligible") else "Pending"
		out[a.name] = {
			"application_id": a.application_id, "candidate_name": a.candidate_name, "email": a.email, "mobile": a.mobile,
			"job_title": a.job_title, "stage": next((st for st, statuses in SUMMARY_STAGES if a.status in statuses), a.status),
			"status": a.status, "eligibility": _("To check") if elig == "Pending" else _(elig), "eligibility_key": elig,
			"reason": a.eligibility_reason, "applied_on": a.application_date,
		}
	return out


def _excel_chart(chart, ws, first_row, last_row, columns):
	"""A native Excel chart of the same kind over the sheet's table."""
	from openpyxl.chart import BarChart, DoughnutChart, LineChart, PieChart, Reference
	from openpyxl.chart.label import DataLabelList
	from openpyxl.chart.series import DataPoint

	kind = chart["kind"]
	labels = Reference(ws, min_col=1, min_row=first_row + 1, max_row=last_row)
	if kind == "stacked":
		xl = BarChart()
		xl.type, xl.grouping, xl.overlap = "bar", "stacked", 100
		xl.add_data(Reference(ws, min_col=2, max_col=columns - 1, min_row=first_row, max_row=last_row), titles_from_data=True)
		for series, s in zip(xl.series, chart["series"]):
			series.graphicalProperties.solidFill = s["color"][1:]
			series.graphicalProperties.line.solidFill = "FFFFFF"
	else:
		if kind in ("pie", "donut"):
			xl = DoughnutChart(holeSize=55) if kind == "donut" else PieChart()
		elif kind == "line":
			xl = LineChart()
			xl.legend = None
		else:
			xl = BarChart()
			xl.type = "col" if kind == "column" else "bar"
			xl.legend = None
			xl.y_axis.majorGridlines = None
		xl.add_data(Reference(ws, min_col=2, min_row=first_row, max_row=last_row), titles_from_data=True)
		series = xl.series[0]
		if kind == "line":
			series.graphicalProperties.line.solidFill = chart["color"][1:]
			series.graphicalProperties.line.width = 22000  # EMU, about 1.75pt
			series.smooth = False
		else:
			for idx, s in enumerate(chart["slices"]):
				point = DataPoint(idx=idx)
				point.graphicalProperties.solidFill = s["color"][1:]
				point.graphicalProperties.line.solidFill = "FFFFFF"
				series.dPt.append(point)
			series.dLbls = DataLabelList()
			series.dLbls.showVal = True
			series.dLbls.numFmt = "0;;;"  # hide labels of empty groups
	xl.set_categories(labels)
	xl.title = chart["title"]
	rows = last_row - first_row
	if kind in ("bar", "funnel", "stacked"):
		xl.x_axis.scaling.orientation = "maxMin"  # first row at the top, as on the page
		xl.height, xl.width = max(7.5, 0.6 * rows + 3), 22
	elif kind == "line":
		xl.height, xl.width = 8, max(18, min(40, rows * 0.6))
	else:
		xl.height, xl.width = 9, 15
	return xl


def _workbook(data, applied):
	from openpyxl import Workbook
	from openpyxl.styles import Font, PatternFill
	from openpyxl.utils import get_column_letter
	from openpyxl.worksheet.hyperlink import Hyperlink

	from pathways.pathways.report.recruitment_pipeline.recruitment_pipeline import get_columns, get_data
	from pathways.utils.application_export import _sheet_name

	header_font = Font(bold=True, color="FFFFFF")
	header_fill = PatternFill("solid", fgColor=_theme().primary[1:].upper())

	def header(ws, row, values):
		for c, value in enumerate(values, start=1):
			cell = ws.cell(row=row, column=c)
			text(cell, value)
			cell.font, cell.fill = header_font, header_fill

	def text(cell, value):
		# Names come from masters and job titles: never a formula.
		cell.value = value
		cell.data_type = "s"

	wb = Workbook()
	summary = wb.active
	summary.title = _("Summary")
	summary.append([_("Recruitment Report")])
	summary["A1"].font = Font(bold=True, size=14)
	summary.append([_("Exported On"), now_datetime()])
	summary["B2"].number_format = "DD-MMM-YYYY HH:MM"
	summary.append([_("Exported By"), frappe.utils.get_fullname(frappe.session.user)])
	for label, value in applied["rows"]:
		summary.append([label, None])
		cell = summary.cell(row=summary.max_row, column=2)
		if isinstance(value, datetime.date):
			cell.value, cell.number_format = value, "DD-MMM-YYYY"
		else:
			text(cell, value)
	summary.append([_("Total Applications"), data["total"]])
	summary.append([])
	header(summary, summary.max_row + 1, [_("Sheet"), _("Chart")])

	used = {summary.title.lower()}
	for chart in data["charts"]:
		name = _sheet_name(chart["title"], used)
		ws = wb.create_sheet(name)
		row = summary.max_row + 1
		escaped = name.replace("'", "''")
		link = summary.cell(row=row, column=1)
		text(link, name)
		link.hyperlink = Hyperlink(ref=f"A{row}", location=f"'{escaped}'!A1")
		link.font = Font(color="0563C1", underline="single")
		text(summary.cell(row=row, column=2), chart["subtitle"])

		ws["A1"] = chart["title"]
		ws["A1"].font = Font(bold=True, size=13)
		text(ws["A2"], chart["subtitle"])
		ws["A2"].font = Font(italic=True, color="52514E")
		table = chart["table"]
		header(ws, 4, table["header"])
		for r, values in enumerate(table["rows"], start=5):
			for c, value in enumerate(values, start=1):
				cell = ws.cell(row=r, column=c)
				if isinstance(value, str):
					text(cell, value)
				else:
					cell.value = value
					if c - 1 in table["percent"]:
						cell.number_format = "0.0%"
		last = 4 + len(table["rows"])
		widest = max([len(str(v[0])) for v in table["rows"]] + [len(table["header"][0])])
		ws.column_dimensions["A"].width = max(16, min(60, widest + 2))
		for c in range(2, len(table["header"]) + 1):
			ws.column_dimensions[get_column_letter(c)].width = max(12, len(str(table["header"][c - 1])) + 2)
		ws.freeze_panes = "A5"

		if table["rows"] and chart["total"]:
			anchor = get_column_letter(len(table["header"]) + 2) + "4"
			ws.add_chart(_excel_chart(chart, ws, 4, last, len(table["header"])), anchor)

	# The pipeline table, job titles instead of IDs.
	ws = wb.create_sheet(_sheet_name(_("Pipeline by Job"), used))
	columns = get_columns()
	rows = [r for r in get_data({}) if applied.get("jobs") is None or r["job_opening"] in applied["jobs"]]
	header(ws, 1, [_(c["label"]) for c in columns])
	for r, row in enumerate(rows, start=2):
		for c, col in enumerate(columns, start=1):
			value = row.get(col["fieldname"])
			cell = ws.cell(row=r, column=c)
			if col["fieldtype"] == "Int":
				cell.value = value or 0
			else:
				text(cell, value or "")
	for c, col in enumerate(columns, start=1):
		ws.column_dimensions[get_column_letter(c)].width = 44 if col["fieldname"] == "job_title" else 14
	ws.freeze_panes = "A2"

	summary.column_dimensions["A"].width = 30
	summary.column_dimensions["B"].width = 60
	return wb
