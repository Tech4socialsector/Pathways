# Copyright (c) 2026, NLSIU and contributors
# For license information, please see license.txt

"""Shortlisting sheets (workflow folder "4. Shortlisting"), one layout per
track, built from the applications and the committee's decisions:

- Admin:    Shortlisting Sheet (candidate details, screening answers,
            Ineligible + reason, Shortlisted + remarks)
- Research: Shortlisting Sheet (Assessed by ... rubric scores, Total, Status)
- Faculty:  Eligibility Check & Scoring, and Post Eligibility Check
            (Consolidated scores of eligible candidates)

get_shortlisting_report previews them; download_shortlisting_report
returns the same sheets as an .xlsx file."""

from io import BytesIO

import frappe
from frappe import _
from frappe.utils import cint, date_diff, flt, get_url, getdate, today

from pathways.api.scoring import can_shortlist

PAST_SHORTLISTING = {
	"Shortlisted", "Interview Scheduled", "Interview Completed", "Selected", "Offer Extended",
	"Offer Accepted", "Offer Declined", "Documents Pending", "Documents Verified", "Joined",
}


# ------------------------------------------------------------------ helpers


def _age(dob):
	return int(date_diff(today(), getdate(dob)) / 365.25) if dob else ""


def _link(url):
	return get_url(url) if url else ""


def _yes_no(flag):
	return "Yes" if flag else "No"


def _qual(app, level):
	rows = [q for q in app.qualifications or [] if q.degree_level == level]
	return rows[0] if rows else frappe._dict()


def _score_for(app_name):
	name = frappe.db.get_value("Shortlisting Score", {"application": app_name}, "name", order_by="modified desc")
	if not name:
		return None
	score = frappe.get_doc("Shortlisting Score", name)
	score.by_label = {c.criterion_label: c.score_given for c in score.criteria or []}
	return score


def _decision(app):
	if app.status in PAST_SHORTLISTING:
		return "Yes"
	if app.status == "Not Selected" and app.eligibility_status != "Not Eligible":
		return "No"
	return ""


def _document(app, *words):
	for d in app.documents or []:
		if all(w.lower() in (d.document_type or "").lower() for w in words):
			return _link(d.attachment)
	return ""


def _latest_job(app):
	rows = sorted(app.employment_history or [], key=lambda r: (cint(r.is_current), str(r.from_date or "")), reverse=True)
	return rows[0] if rows else frappe._dict()


def _screening(job, apps):
	"""One column per screening question (and its details), in form order."""
	questions = [q.question for q in job.get("screening_questions") or []]
	for app in apps:
		for a in app.screening_answers or []:
			if a.question not in questions:
				questions.append(a.question)
	cols = []
	for q in questions:
		cols.append(q)
		if any(a.details for app in apps for a in app.screening_answers or [] if a.question == q):
			cols.append(_("If yes, details: {0}").format(q[:60]))
	return questions, cols


def _screening_values(app, questions, cols):
	answers = {a.question: a for a in app.screening_answers or []}
	out = []
	for q in questions:
		a = answers.get(q) or frappe._dict()
		out.append(a.answer or "")
		if _("If yes, details: {0}").format(q[:60]) in cols:
			out.append(a.details or "")
	return out


def _criteria(job):
	from pathways.api.scoring import get_rubric_for_job_opening

	rubric = get_rubric_for_job_opening(job.name, "Shortlisting") or {}
	scored = rubric.get("scoring_mode") != "Pass/Fail Only"
	return [c for c in rubric.get("criteria") or []] if scored else []


def _criterion_header(c):
	return f"{c['criterion_label']} (MAX SCORE = {flt(c['max_score']):g})"


# ------------------------------------------------------------------ layouts


def _admin(job, apps, cands):
	questions, q_cols = _screening(job, apps)
	columns = [
		"Sl. No", "Application ID", "Email Address", "Name", "Age", "Gender", "Mobile Number",
		"Overall work experience (Number of Years)", "Relevant work experience (Number of Years)",
		"Name of the Degree", "College / University", "Year of Graduation", "Percentage Secured / CGPA Secured",
		"Division / Grade", "Specialization",
		"Designation", "Name of the employer", "From Date", "To Date", "Key Responsibilities", "Current Notice Period",
		*q_cols,
		"Ineligible", "Reason for Ineligibility", "Shortlisted",
		"Remarks/Comments (Stating the reason for shortlisting & not shortlisting)",
	]
	rows = []
	for i, app in enumerate(apps, 1):
		c = cands[app.candidate]
		q = _qual(app, "Postgraduate") or _qual(app, "Undergraduate")
		j = _latest_job(app)
		score = _score_for(app.name)
		rows.append([
			i, app.application_id, c.email, c.full_name, _age(c.date_of_birth), c.gender, c.mobile_number,
			app.overall_experience_years, app.relevant_experience_years,
			q.degree_name, q.institution or q.other_institution, q.year_of_graduation, q.percentage_or_cgpa,
			q.division_grade, q.specialization,
			j.designation, j.employer_name, j.from_date, "Present" if j.is_current else j.to_date, j.key_responsibilities,
			app.notice_period or j.notice_period,
			*_screening_values(app, questions, q_cols),
			"Yes" if app.eligibility_status == "Not Eligible" else ("No" if app.eligibility_status == "Eligible" else ""),
			app.eligibility_reason, _decision(app), (score.remarks if score else ""),
		])
	return [{"title": "Shortlisting Sheet", "columns": columns, "rows": rows}]


def _research(job, apps, cands):
	questions, q_cols = _screening(job, apps)
	criteria = _criteria(job)
	max_total = sum(flt(c["max_score"]) for c in criteria)
	columns = [
		"Assessed by", "Application ID", "Email Id", "Name", "Age", "Gender",
		"Overall work experience", "Relevant work experience (Number of Years)",
		*q_cols, "Eligible?",
		"Name of the Degree (Bachelor's Degree)", "Applicant's area(s) of specialization", "Percentage",
		"Name of the Degree (Master's Degree)", "Applicant's area(s) of specialization (Master's)", "Percentage (Master's)",
		"Resume / CV", "Statement of purpose", "Writing sample", "Research publication / equivalent output",
		*[_criterion_header(c) for c in criteria],
		f"TOTAL SCORE (BEST SCORE = {max_total:g})", "Remarks/Comments", "Current Notice Period (if applicable)", "Status",
	]
	rows = []
	for app in apps:
		c = cands[app.candidate]
		ug, pg = _qual(app, "Undergraduate"), _qual(app, "Postgraduate")
		score = _score_for(app.name)
		pub = (app.publications or [frappe._dict()])[0]
		rows.append([
			frappe.utils.get_fullname(score.owner) if score else "", app.application_id, c.email, c.full_name,
			_age(c.date_of_birth), c.gender, app.overall_experience_years, app.relevant_experience_years,
			*_screening_values(app, questions, q_cols), app.eligibility_status or "Pending",
			ug.degree_name, ug.specialization, ug.percentage_or_cgpa,
			pg.degree_name, pg.specialization, pg.percentage_or_cgpa,
			_link(app.resume_attachment), _link(app.sop_attachment), _document(app, "writing"),
			_link(pub.pdf_attachment) or pub.doi_link or _document(app, "publication"),
			*[(score.by_label.get(cr["criterion_label"], "") if score else "") for cr in criteria],
			(score.total_score if score else ""), (score.remarks if score else ""), app.notice_period, app.status,
		])
	return [{"title": "Shortlisting Sheet", "columns": columns, "rows": rows}]


def _faculty(job, apps, cands):
	criteria = _criteria(job)
	max_total = sum(flt(c["max_score"]) for c in criteria)
	eligibility = {
		"title": "Eligibility Check",
		"columns": [
			"S No.", "Application ID", "Name", "Email Address", "Mobile Number", "Date of Birth", "Age", "Gender", "Category",
			"Applicant's Area of Specialization", "Other Specialization", "Which position are you applying for?",
			"Master's Degree Percentage / CGPA", "Have you been awarded your PhD degree?",
			"Have you successfully cleared the NET/ SLET/ SET?", "NET Exam", "NET Subject",
			"Type of Physical Disability", "Percentage of Disability",
			"Undergraduate Degree", "UG College / University", "UG Year of Graduation", "UG Percentage", "UG CGPA (scale)",
			"Postgraduate Degree", "PG College / University", "PG Year of Graduation", "PG Percentage", "PG CGPA (scale)",
			"PhD University", "PhD QS / THE / ARWU rank",
			"Overall work experience in MONTHS", "Teaching experience in MONTHS", "Research experience in MONTHS",
			"Professional Legal experience in MONTHS", "Number of publications",
			"Eligibility (yes/no)", "Reason for ineligibility",
		],
		"rows": [],
	}
	consolidated = {
		"title": "Consolidated scores",
		"columns": [
			"S No.", "Application ID", "Eligible Candidate Name", "Subject", "Applied For", "Area of Specialisation",
			"Phone Number", "Email", *[_criterion_header(c) for c in criteria],
			f"TOTAL SCORE (MAX SCORE = {max_total:g})", "Remarks",
			"Publication #1 - Name of the Journal", "Publication #1 - Link of the article - DOI", "Publication #1 - PDF",
			"Resume", "Shortlisted",
		],
		"rows": [],
	}
	for i, app in enumerate(apps, 1):
		c = cands[app.candidate]
		ug, pg, phd = _qual(app, "Undergraduate"), _qual(app, "Postgraduate"), _qual(app, "Doctoral")
		cgpa = lambda q: f"{q.cgpa} ({q.cgpa_scale})" if q.cgpa else ""
		eligibility["rows"].append([
			i, app.application_id, c.full_name, c.email, c.mobile_number, c.date_of_birth, _age(c.date_of_birth), c.gender,
			c.get("category"), (app.specializations or "").replace("\n", ", "), app.other_specialization, job.job_title,
			pg.percentage_or_cgpa, app.phd_awarded, app.net_qualified, app.net_exam, app.net_subject or app.net_other_subject,
			app.disability_type, app.disability_percentage,
			ug.degree_name, ug.institution or ug.other_institution, ug.year_of_graduation, ug.percentage_or_cgpa, cgpa(ug),
			pg.degree_name, pg.institution or pg.other_institution, pg.year_of_graduation, pg.percentage_or_cgpa, cgpa(pg),
			phd.institution or phd.other_institution, " / ".join(str(r or "-") for r in (phd.qs_rank, phd.the_rank, phd.arwu_rank)) if phd else "",
			app.overall_experience_months, app.teaching_experience_months, app.research_experience_months,
			app.legal_experience_months, len(app.publications or []),
			{"Eligible": "Yes", "Not Eligible": "No"}.get(app.eligibility_status, "Pending"), app.eligibility_reason,
		])
		if app.eligibility_status != "Eligible":
			continue
		score = _score_for(app.name)
		pub = (app.publications or [frappe._dict()])[0]
		consolidated["rows"].append([
			len(consolidated["rows"]) + 1, app.application_id, c.full_name, job.get("specialization_discipline") or "",
			job.job_title, (app.specializations or "").replace("\n", ", "), c.mobile_number, c.email,
			*[(score.by_label.get(cr["criterion_label"], "") if score else "") for cr in criteria],
			(score.total_score if score else ""), (score.remarks if score else ""),
			pub.journal_name, pub.doi_link, _link(pub.pdf_attachment), _link(app.resume_attachment), _decision(app),
		])
	return [eligibility, consolidated]


LAYOUTS = {"Admin": _admin, "Research": _research, "Faculty": _faculty}


def build_report(job_opening):
	job = frappe.get_doc("Job Opening", job_opening)
	if not can_shortlist(job.name):
		frappe.throw(_("Only the shortlisting committee for this job can see its shortlisting sheet."), frappe.PermissionError)
	names = frappe.get_all(
		"Application", filters={"job_opening": job.name, "status": ["!=", "Withdrawn"]}, pluck="name", order_by="creation asc"
	)
	apps = [frappe.get_doc("Application", n) for n in names]
	cand_fields = ["name", "full_name", "email", "mobile_number", "date_of_birth", "gender"]
	if frappe.get_meta("Candidate").has_field("category"):
		cand_fields.append("category")
	cands = {
		c.name: c
		for c in frappe.get_all("Candidate", filters={"name": ["in", [a.candidate for a in apps] or [""]]}, fields=cand_fields)
	}
	layout = LAYOUTS.get(job.track, _admin)
	return {"job_title": job.job_title, "job_code": job.position or job.name, "track": job.track, "sheets": layout(job, apps, cands)}


@frappe.whitelist()
def get_shortlisting_report(job_opening):
	report = build_report(job_opening)
	for sheet in report["sheets"]:
		sheet["rows"] = [[_plain(v) for v in row] for row in sheet["rows"]]
	return report


def _plain(value):
	if value is None:
		return ""
	if hasattr(value, "isoformat"):
		return frappe.utils.formatdate(value) if not hasattr(value, "hour") else frappe.utils.format_datetime(value)
	return value


@frappe.whitelist()
def download_shortlisting_report(job_opening):
	report = build_report(job_opening)
	safe = "".join(ch if ch.isalnum() or ch in " -_" else "_" for ch in report["job_title"]).strip()
	_send_workbook(report["sheets"], f"Shortlisting Sheet - {safe}.xlsx")


def _send_workbook(sheets, filename):
	from openpyxl import Workbook
	from openpyxl.styles import Alignment, Font, PatternFill
	from openpyxl.utils import get_column_letter

	wb = Workbook()
	wb.remove(wb.active)
	header_fill = PatternFill("solid", fgColor="920C24")
	for sheet in sheets:
		ws = wb.create_sheet(sheet["title"][:31])
		ws.append(sheet["columns"])
		for row in sheet["rows"]:
			ws.append([_plain(v) if hasattr(v, "isoformat") else v for v in row])
		for idx, label in enumerate(sheet["columns"], 1):
			cell = ws.cell(row=1, column=idx)
			cell.font = Font(bold=True, color="FFFFFF")
			cell.fill = header_fill
			cell.alignment = Alignment(wrap_text=True, vertical="top")
			ws.column_dimensions[get_column_letter(idx)].width = min(max(len(str(label)) * 0.6, 12), 40)
		ws.freeze_panes = "D2"
		ws.row_dimensions[1].height = 60
	out = BytesIO()
	wb.save(out)
	frappe.response["filename"] = filename
	frappe.response["filecontent"] = out.getvalue()
	frappe.response["type"] = "binary"


# ------------------------------------------------------------- all jobs


def _jobs_for_reports(track=None, job_opening=None, job_openings=None):
	"""Job openings with applications that the user may shortlist."""
	filters = {"name": job_opening} if job_opening else {}
	if job_openings:
		names = frappe.parse_json(job_openings) if isinstance(job_openings, str) else job_openings
		filters["name"] = ["in", list(names) or [""]]
	if track:
		filters["track"] = track
	jobs = frappe.get_all(
		"Job Opening",
		filters=filters,
		fields=["name", "job_title", "position", "track", "department", "status", "vacancies", "shortlisting_ratio"],
		order_by="creation desc",
	)
	return [j for j in jobs if frappe.db.exists("Application", {"job_opening": j.name}) and can_shortlist(j.name)]


@frappe.whitelist()
def list_shortlisting_sheets(track=None, job_opening=None, job_openings=None):
	"""Reports page: each job's shortlisting position at a glance."""
	from pathways.api.scoring import shortlisting_summary

	out = []
	for job in _jobs_for_reports(track, job_opening, job_openings):
		s = shortlisting_summary(frappe.get_doc("Job Opening", job.name))
		out.append(
			{
				"job_opening": job.name,
				"job_title": job.job_title,
				"job_code": job.position or job.name,
				"track": job.track,
				"department": job.department,
				"status": job.status,
				"applications": s["applications"],
				"eligible": s["eligible"],
				"not_eligible": s["not_eligible"],
				"pending": s["pending"],
				"scored": s["scored"],
				"shortlisted": s["shortlisted"],
				"target": s["target"],
				"committee": ", ".join(m["full_name"] or m["user"] for m in (s["committee"] or {}).get("members", [])),
			}
		)
	return out


def _overall_sheet(jobs):
	columns = [
		"Job Code", "Job Opening", "Track", "Department", "Application ID", "Name", "Email", "Mobile Number",
		"Applied On", "Eligibility", "Reason for Ineligibility", "Status", "Total Score", "Best Score",
		"Shortlisted", "Remarks", "Assessed by", "Regret sent on",
	]
	rows = []
	for job in jobs:
		criteria = _criteria(frappe.get_doc("Job Opening", job.name))
		best = sum(flt(c["max_score"]) for c in criteria) or ""
		for app in frappe.get_all(
			"Application",
			filters={"job_opening": job.name, "status": ["!=", "Withdrawn"]},
			fields=["name", "application_id", "candidate", "application_date", "eligibility_status", "eligibility_reason", "status", "regret_sent_on"],
			order_by="creation asc",
		):
			c = frappe.db.get_value("Candidate", app.candidate, ["full_name", "email", "mobile_number"], as_dict=True) or {}
			score = _score_for(app.name)
			rows.append([
				job.position or job.name, job.job_title, job.track, job.department, app.application_id, c.get("full_name"),
				c.get("email"), c.get("mobile_number"), app.application_date, app.eligibility_status or "Pending",
				app.eligibility_reason, app.status, (score.total_score if score and criteria else ""), best,
				_decision(app), (score.remarks if score else ""), (frappe.utils.get_fullname(score.owner) if score else ""),
				app.regret_sent_on,
			])
	return {"title": "Overall", "columns": columns, "rows": rows}


@frappe.whitelist()
def download_all_shortlisting_sheets(track=None, job_opening=None, job_openings=None):
	"""One workbook: an Overall sheet with every candidate, then each job's
	own shortlisting sheet(s)."""
	jobs = _jobs_for_reports(track, job_opening, job_openings)
	if not jobs:
		frappe.throw(_("No job openings with applications match these filters."))
	sheets = [_overall_sheet(jobs)]
	used = {"Overall"}
	for job in jobs:
		for sheet in build_report(job.name)["sheets"]:
			title = f"{job.position or job.name} {sheet['title']}"[:31]
			n = 2
			while title in used:
				title = f"{title[:28]} {n}"
				n += 1
			used.add(title)
			sheets.append({**sheet, "title": title})
	_send_workbook(sheets, "Shortlisting Sheets - All Jobs.xlsx")
